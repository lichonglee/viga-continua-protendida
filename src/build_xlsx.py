"""Gera a planilha de apoio (formulas vivas) - viga continua protendida 2 x 10 m."""
import sys
sys.path.insert(0, '/tmp/claude-0/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/scratchpad')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from analysis import final_design

d = final_design()
NAVY = '1F3864'; LBLUE = 'DCE6F1'; YEL = 'FFF2CC'; GRY = 'F2F2F2'; GRN = 'E2EFDA'
fh = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
fb = Font(name='Calibri', size=10)
fbb = Font(name='Calibri', size=10, bold=True)
fill_h = PatternFill('solid', fgColor=NAVY)
fill_in = PatternFill('solid', fgColor=YEL)
fill_res = PatternFill('solid', fgColor=LBLUE)
fill_sec = PatternFill('solid', fgColor='D9D9D9')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
FIRST = 8            # primeira linha de estacoes
NST = 41
LAST = FIRST + NST - 1
DX = 0.5

# ------------------------------------------------------------------ Dados
wsD = wb.active; wsD.title = 'Dados'
rows = []   # (key, label, value, unit, note, kind)   kind: 'in' | 'f' | 'sec'
def sec(t): rows.append((None, t, None, None, None, 'sec'))
def inp(k, lab, v, u='', note=''): rows.append((k, lab, v, u, note, 'in'))
def frm(k, lab, f, u='', note=''): rows.append((k, lab, f, u, note, 'f'))

sec('1. Geometria e materiais')
inp('L', 'Vão de cada tramo', 10.0, 'm', 'viga contínua de 2 vãos iguais')
inp('bw', 'Largura da seção', 0.30, 'm')
inp('h', 'Altura da seção', 0.90, 'm')
frm('Ac', 'Área Ac', '={bw}*{h}', 'm²')
frm('Ic', 'Inércia Ic', '={bw}*{h}^3/12', 'm⁴')
frm('W', 'Módulo resistente W1 = W2', '={Ic}/({h}/2)', 'm³')
frm('k', 'Núcleo central k = W/Ac', '={W}/{Ac}', 'm')
inp('fck', 'fck', 30.0, 'MPa')
inp('gc', 'γc', 1.4, '')
frm('fcd', 'fcd', '={fck}/{gc}', 'MPa')
frm('fctm', 'fct,m = 0,3·fck^(2/3)', '=0.3*{fck}^(2/3)', 'MPa')
frm('fctkinf', 'fctk,inf = 0,7·fct,m', '=0.7*{fctm}', 'MPa')
frm('fctd', 'fctd = fctk,inf/γc', '={fctkinf}/{gc}', 'MPa')
frm('Eci', 'Eci = 5600·√fck', '=5600*SQRT({fck})', 'MPa')
frm('ai', 'αi = 0,8+0,2·fck/80 ≤ 1', '=MIN(0.8+0.2*{fck}/80,1)', '')
frm('Ecs', 'Ecs = αi·Eci', '={ai}*{Eci}', 'MPa')
inp('t0', 'Idade da protensão (ato)', 7, 'dias')
frm('fcj', 'fcj = exp[0,38(1-√(28/t))]·fck', '=EXP(0.38*(1-SQRT(28/{t0})))*{fck}', 'MPa')
frm('Ecij', 'Eci,j = 5600·√fcj', '=5600*SQRT({fcj})', 'MPa')
frm('fctmj', 'fct,m,j', '=0.3*{fcj}^(2/3)', 'MPa')
frm('ft_els', 'Limite de tração ELS-F: 1,2·fctk,inf', '=1.2*{fctkinf}', 'MPa')
frm('ft_ato', 'Limite de tração no ato: 1,2·fct,m,j', '=1.2*{fctmj}', 'MPa')
frm('fc_ato', 'Limite de compressão no ato: 0,7·fcj', '=0.7*{fcj}', 'MPa')
frm('fc_els', 'Limite de compressão em serviço: 0,7·fck', '=0.7*{fck}', 'MPa')
sec('2. Ações')
frm('gpp', 'Peso próprio g0 = 25·Ac', '=25*{Ac}', 'kN/m')
inp('gadd', 'Carga permanente adicional (tabuleiro, piso, guarda-corpo)', 13.25, 'kN/m', 'ADOTADO - provisório (confirmar com a biapoiada)')
frm('g', 'Carga permanente total g', '={gpp}+{gadd}', 'kN/m')
inp('q', 'Sobrecarga q (passarela)', 25.0, 'kN/m', 'ADOTADO - provisório (5 kN/m² × 5 m)')
inp('psi1', 'ψ1 (passarela, NBR 8681)', 0.4, '')
inp('psi2', 'ψ2 (passarela, NBR 8681)', 0.3, '')
frm('pbal', 'a) Carga a equilibrar = g + ψ2·q', '={g}+{psi2}*{q}', 'kN/m')
sec('3. Cobrimento, excentricidades e traçado do cabo resultante')
inp('cnom', 'Cobrimento nominal (CAA II, protendido)', 0.035, 'm')
inp('phit', 'Estribo Ø', 0.010, 'm')
inp('dbain', 'Bainha: diâmetro externo (5 Ø12,7)', 0.061, 'm')
frm('dpp', "d' = c + Øestribo + Øbainha/2", '={cnom}+{phit}+{dbain}/2', 'm')
frm('emax', 'b) e_max = h/2 - d\'', '={h}/2-{dpp}', 'm')
inp('e_end', 'Excentricidade na ancoragem', 0.0, 'm', '+ abaixo do CG')
inp('e_span', 'Excentricidade máx. no vão (abaixo do CG)', 0.28, 'm')
inp('e_sup', 'Excentricidade no apoio central (acima do CG)', 0.28, 'm')
inp('lam', 'Comprimento da parábola reversa junto ao apoio', 1.0, 'm')
frm('xi', 'x_i = L - λ (início da parábola reversa)', '={L}-{lam}', 'm')
frm('cc', 'c = e_vão - e_ancoragem', '={e_span}-{e_end}', 'm')
frm('qA', 'coef. A da equação de x0', '={e_sup}+{e_end}', '')
frm('qB', 'coef. B da equação de x0', '={cc}*({xi}+{L})', '')
frm('qC', 'coef. C da equação de x0', '=-{cc}*{xi}*{L}', '')
frm('x0', 'x0 (vértice da parábola do vão)', '=(-{qB}+SQRT({qB}^2-4*{qA}*{qC}))/(2*{qA})', 'm')
frm('aA', 'aA = c/x0²', '={cc}/{x0}^2', '1/m')
frm('aB', 'aB = aA·(xi-x0)/λ', '={aA}*({xi}-{x0})/{lam}', '1/m')
sec('4. Armadura ativa')
inp('ncab', 'Número de cabos', 2, '')
inp('ncord', 'Cordoalhas por cabo', 5, '')
inp('Ap1', 'Área de 1 cordoalha CP190 RB Ø12,7', 0.987, 'cm²')
frm('Ap', 'Ap total', '={ncab}*{ncord}*{Ap1}/10000', 'm²')
inp('fptk', 'fptk (CP190 RB)', 1900.0, 'MPa')
frm('fpyk', 'fpyk = 0,9·fptk (RB)', '=0.9*{fptk}', 'MPa')
inp('gs', 'γs', 1.15, '')
frm('fpyd', 'fpyd', '={fpyk}/{gs}', 'MPa')
frm('fptd', 'fptd', '={fptk}/{gs}', 'MPa')
inp('Ep', 'Ep', 200000.0, 'MPa')
inp('rj', 'σp,j / fptk no macaco (pós-tração, RB: 0,74)', 0.74, '')
frm('sj', 'σp,j', '={rj}*{fptk}', 'MPa')
frm('Pj', 'Força de macaco Pj = Ap·σp,j', '={Ap}*{sj}*1000', 'kN')
sec('5. Perdas de protensão')
inp('mu', 'Coeficiente de atrito μ (cordoalha em bainha metálica)', 0.20, '1/rad')
inp('kk', 'Perda parasita k (=0,01·μ)', 0.002, '1/m')
inp('slip', 'Cravação das cunhas (acomodação)', 0.006, 'm')
inp('phi', 'Coeficiente de fluência φ(t∞,t0)', 2.9, '', 'Tab. 8.2 NBR 6118 (UR≈75%, hfic≈22 cm, t0≈7 d) - interpolado; CONFERIR')
inp('ecs', 'Retração εcs(t∞,t0)', -0.00023, '', 'idem')
frm('ap_ato', 'αp = Ep/Eci,j (ato)', '={Ep}/{Ecij}', '')
frm('ap28', 'αp = Ep/Ecs (progressivas)', '={Ep}/{Ecs}', '')
frm('cES', 'Fator de encurtamento elástico (n-1)/(2n)·αp', '={ap_ato}*({ncab}-1)/(2*{ncab})', '')
frm('rho_p', 'ρp = Ap/Ac', '={Ap}/{Ac}', '')
# resultados calculados na aba Tracado_Perdas
frm('kc', 'Nº de estações do trecho 0..L com Δ < cravação', f'=COUNTIF(Tracado_Perdas!$O${FIRST}:$O${FIRST+20},"<"&{{slip}})', '', 'auxiliar - zona de acomodação')
frm('fint', 'Fração de interpolação', f'=({{slip}}-INDEX(Tracado_Perdas!$O${FIRST}:$O${FIRST+20},{{kc}}))/(INDEX(Tracado_Perdas!$O${FIRST}:$O${FIRST+20},{{kc}}+1)-INDEX(Tracado_Perdas!$O${FIRST}:$O${FIRST+20},{{kc}}))', '')
frm('xa', 'Alcance da acomodação xa (a partir de cada extremidade)', f'=INDEX(Tracado_Perdas!$B${FIRST}:$B${FIRST+20},{{kc}})+{{fint}}*{DX}', 'm')
frm('sfa', 'σ(xa) pelo atrito', f'=INDEX(Tracado_Perdas!$M${FIRST}:$M${FIRST+20},{{kc}})+{{fint}}*(INDEX(Tracado_Perdas!$M${FIRST}:$M${FIRST+20},{{kc}}+1)-INDEX(Tracado_Perdas!$M${FIRST}:$M${FIRST+20},{{kc}}))', 'MPa')
sec('6. Armaduras passivas / cisalhamento')
inp('fyk', 'fyk (CA-50)', 500.0, 'MPa')
frm('fyd', 'fyd', '={fyk}/{gc}*{gc}/{gs}', 'MPa')
inp('Es', 'Es', 210000.0, 'MPa')
frm('dd', 'd = 0,9·h (cisalhamento)', '=0.9*{h}', 'm')
frm('tauc', 'τc = 0,6·fctd', '=0.6*{fctd}', 'MPa')
frm('tauu', 'τu = 0,27·(1-fck/250)·fcd', '=0.27*(1-{fck}/250)*{fcd}', 'MPa')
frm('rhomin', 'ρsw,min = 0,2·fct,m/fywk', '=0.2*{fctm}/{fyk}', '')
frm('bwred', 'bw,red = bw - ½ΣØbainha', '=IF({dbain}>{bw}/8,{bw}-0.5*{ncab}*{dbain},{bw})', 'm')

A = {}   # key -> address
r0 = 4
rr = r0
for (k, lab, v, u, note, kind) in rows:
    if k: A[k] = f'$C${rr}'
    rr += 1
def Dd(k): return 'Dados!' + A[k]

wsD['A1'] = 'PLANILHA DE APOIO - VIGA CONTÍNUA PROTENDIDA (2 × 10 m)'; wsD['A1'].font = Font(bold=True, size=14, color=NAVY)
wsD['A2'] = 'Células amarelas = dados de entrada. Todas as demais abas são recalculadas a partir desta aba. Unidades: kN, m, MPa (salvo indicação).'
wsD['A2'].font = Font(italic=True, size=9)
rr = r0
for (k, lab, v, u, note, kind) in rows:
    if kind == 'sec':
        wsD.cell(rr, 1, lab).font = fbb
        for c in range(1, 6):
            wsD.cell(rr, c).fill = fill_sec
    else:
        wsD.cell(rr, 1, k).font = Font(size=8, color='808080')
        wsD.cell(rr, 2, lab).font = fb
        cell = wsD.cell(rr, 3)
        if kind == 'in':
            cell.value = v; cell.fill = fill_in
        else:
            cell.value = v.format(**{kk: vv.replace('$', '') if False else vv for kk, vv in A.items()}) if isinstance(v, str) else v
            cell.fill = fill_res
        cell.font = fb; cell.border = box
        wsD.cell(rr, 4, u).font = fb
        wsD.cell(rr, 5, note).font = Font(size=9, italic=True, color='C00000' if 'ADOTADO' in note or 'CONFERIR' in note else '595959')
    rr += 1
wsD.column_dimensions['A'].width = 8; wsD.column_dimensions['B'].width = 58
wsD.column_dimensions['C'].width = 14; wsD.column_dimensions['D'].width = 9; wsD.column_dimensions['E'].width = 60
for rrr in range(r0, rr):
    wsD.cell(rrr, 3).number_format = '0.0000'

# helper to make a station sheet ------------------------------------------------
def station_sheet(name, title, cols, widths=None, note=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=13, color=NAVY)
    if note:
        ws['A2'] = note; ws['A2'].font = Font(italic=True, size=9)
    for j, (key, head, unit, fn, fmt) in enumerate(cols, start=1):
        h = ws.cell(6, j, head); h.font = fh; h.fill = fill_h
        h.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        u = ws.cell(7, j, unit); u.font = Font(size=8, italic=True); u.alignment = Alignment(horizontal='center')
        ws.column_dimensions[CL(j)].width = (widths or {}).get(key, 11)
    ws.row_dimensions[6].height = 32
    colmap = {key: CL(j) for j, (key, *_rest) in enumerate(cols, start=1)}
    for i in range(NST):
        r = FIRST + i
        for j, (key, head, unit, fn, fmt) in enumerate(cols, start=1):
            val = fn(r, i, colmap)
            c = ws.cell(r, j, val)
            c.font = fb; c.border = box
            if fmt: c.number_format = fmt
    ws.freeze_panes = ws.cell(FIRST, 3)
    return ws, colmap

def prev(r, col): return f'{col}{r-1}'

# ------------------------------------------------------------------ Esforcos
q_, g_, g0_, L_ = Dd('q'), Dd('g'), Dd('gpp'), Dd('L')
def Mboth(w, s): return f'(3*{w}*{L_}*{s}/8-{w}*{s}^2/2)'
def Mown(w, s): return f'(7*{w}*{L_}*{s}/16-{w}*{s}^2/2)'
def Moth(w, s): return f'(-{w}*{L_}*{s}/16)'
def Vboth(w, s): return f'(3*{w}*{L_}/8-{w}*{s})'
def Vown(w, s): return f'(7*{w}*{L_}/16-{w}*{s})'
def Voth(w, s): return f'(-{w}*{L_}/16)'
cE = [
    ('i', 'i', '', lambda r, i, c: i, '0'),
    ('x', 'x', 'm', lambda r, i, c: f'=A{r}*{DX}', '0.0'),
    ('s', 's (espelho)', 'm', lambda r, i, c: f'=IF(B{r}<={L_},B{r},2*{L_}-B{r})', '0.0'),
    ('Mg0', 'Mg0 (g0)', 'kN·m', lambda r, i, c: f'={Mboth(g0_, "C"+str(r))}', '0.0'),
    ('Mg', 'Mg (g)', 'kN·m', lambda r, i, c: f'={Mboth(g_, "C"+str(r))}', '0.0'),
    ('Mqmax', 'Mq,máx', 'kN·m', lambda r, i, c: f'=MAX(0,{Mboth(q_, "C"+str(r))},{Mown(q_, "C"+str(r))},{Moth(q_, "C"+str(r))})', '0.0'),
    ('Mqmin', 'Mq,mín', 'kN·m', lambda r, i, c: f'=MIN(0,{Mboth(q_, "C"+str(r))},{Mown(q_, "C"+str(r))},{Moth(q_, "C"+str(r))})', '0.0'),
    ('Mcfqmax', 'M CFQ,máx = Mg+ψ1Mq+', 'kN·m', lambda r, i, c: f'=E{r}+{Dd("psi1")}*F{r}', '0.0'),
    ('Mcfqmin', 'M CFQ,mín = Mg+ψ1Mq-', 'kN·m', lambda r, i, c: f'=E{r}+{Dd("psi1")}*G{r}', '0.0'),
    ('Mcrmax', 'M CR,máx = Mg+Mq+', 'kN·m', lambda r, i, c: f'=E{r}+F{r}', '0.0'),
    ('Mcrmin', 'M CR,mín = Mg+Mq-', 'kN·m', lambda r, i, c: f'=E{r}+G{r}', '0.0'),
    ('Vg', 'Vg', 'kN', lambda r, i, c: f'=IF(B{r}<={L_},1,-1)*{Vboth(g_, "C"+str(r))}', '0.0'),
    ('Vqmax', 'Vq,máx', 'kN', lambda r, i, c: f'=IF(B{r}<={L_},MAX(0,{Vboth(q_, "C"+str(r))},{Vown(q_, "C"+str(r))},{Voth(q_, "C"+str(r))}),-MIN(0,{Vboth(q_, "C"+str(r))},{Vown(q_, "C"+str(r))},{Voth(q_, "C"+str(r))}))', '0.0'),
    ('Vqmin', 'Vq,mín', 'kN', lambda r, i, c: f'=IF(B{r}<={L_},MIN(0,{Vboth(q_, "C"+str(r))},{Vown(q_, "C"+str(r))},{Voth(q_, "C"+str(r))}),-MAX(0,{Vboth(q_, "C"+str(r))},{Vown(q_, "C"+str(r))},{Voth(q_, "C"+str(r))}))', '0.0'),
]
wsE, cEm = station_sheet('Esforcos', 'Esforços externos - viga contínua de 2 vãos iguais (envoltórias)', cE,
                         note='Momento positivo = tração embaixo. Padrões de sobrecarga: ambos os vãos, só o vão do ponto, só o outro vão. Estações a cada 0,5 m (0 a 20 m).')

# ------------------------------------------------------------------ Tracado_Perdas
Dx = Dd
def TP(key): return None
cT = [
    ('i', 'i', '', lambda r, i, c: i, '0'),
    ('x', 'x', 'm', lambda r, i, c: f'=A{r}*{DX}', '0.0'),
    ('s', 's', 'm', lambda r, i, c: f'=IF(B{r}<={L_},B{r},2*{L_}-B{r})', '0.0'),
    ('e', 'e(x) (+ abaixo do CG)', 'm', lambda r, i, c: f'=IF(C{r}<={Dd("xi")},{Dd("e_span")}-{Dd("aA")}*(C{r}-{Dd("x0")})^2,-{Dd("e_sup")}+{Dd("aB")}*({L_}-C{r})^2)', '0.0000'),
    ('de', 'de/dx', '', lambda r, i, c: f'=IF(B{r}<={L_},1,-1)*IF(C{r}<={Dd("xi")},-2*{Dd("aA")}*(C{r}-{Dd("x0")}),-2*{Dd("aB")}*({L_}-C{r}))', '0.00000'),
    ('al', 'α', 'rad', lambda r, i, c: f'=ATAN(E{r})', '0.00000'),
    ('ca', 'cos α', '', lambda r, i, c: f'=COS(F{r})', '0.00000'),
    ('sa', 'sen α', '', lambda r, i, c: f'=SIN(F{r})', '0.00000'),
    ('dth', 'Δα', 'rad', lambda r, i, c: 0 if i == 0 else f'=ABS(F{r}-F{r-1})', '0.00000'),
    ('th', 'Σα', 'rad', lambda r, i, c: f'=I{r}' if i == 0 else f'=J{r-1}+I{r}', '0.00000'),
    ('dar', 'Δs (arco)', 'm', lambda r, i, c: 0 if i == 0 else f'=SQRT((B{r}-B{r-1})^2+(D{r}-D{r-1})^2)', '0.0000'),
    ('ar', 'Σs', 'm', lambda r, i, c: f'=K{r}' if i == 0 else f'=L{r-1}+K{r}', '0.000'),
    ('sL', 'σp atrito (da extremidade esq.)', 'MPa', lambda r, i, c: f'={Dd("sj")}*EXP(-{Dd("mu")}*J{r}-{Dd("kk")}*L{r})', '0.0'),
    ('cum', '∫σ dx', 'MPa·m', lambda r, i, c: 0 if i == 0 else f'=N{r-1}+0.5*(M{r}+M{r-1})*(B{r}-B{r-1})', '0.00'),
    ('dk', 'Δ(x) = 2[∫σ - σ·x]/Ep', 'm', lambda r, i, c: f'=2*(N{r}-M{r}*B{r})/{Dd("Ep")}' if i <= 20 else '', '0.00000'),
    ('ss', 'σp após cravação', 'MPa', lambda r, i, c: (f'=IF(B{r}<={Dd("xa")},2*{Dd("sfa")}-M{r},M{r})' if i <= 20 else f'=P{FIRST+40-i}'), '0.0'),
    ('des', 'Δσ enc. elástico', 'MPa', lambda r, i, c: f'=MAX({Dd("cES")}*({Dd("Ap")}*(1/{Dd("Ac")}+D{r}^2/{Dd("Ic")})*P{r}-Esforcos!D{r}*D{r}/{Dd("Ic")}/1000)/(1+{Dd("cES")}*{Dd("Ap")}*(1/{Dd("Ac")}+D{r}^2/{Dd("Ic")})),0)', '0.00'),
    ('si', 'σp,i (ato)', 'MPa', lambda r, i, c: f'=P{r}-Q{r}', '0.0'),
    ('Pi', 'Pi (ato)', 'kN', lambda r, i, c: f'={Dd("Ap")}*R{r}*1000', '0.0'),
    ('scp', 'σcp0 (concreto no cabo)', 'MPa', lambda r, i, c: f'=(-S{r}/{Dd("Ac")}-S{r}*D{r}^2/{Dd("Ic")}+Esforcos!D{r}*D{r}/{Dd("Ic")})/1000', '0.000'),
    ('rr', 'σp,i/fptk', '', lambda r, i, c: f'=R{r}/{Dd("fptk")}', '0.0000'),
    ('psi1000', 'ψ1000', '', lambda r, i, c: f'=IF(U{r}<=0.5,0,IF(U{r}<=0.6,(U{r}-0.5)/0.1*1.3,IF(U{r}<=0.7,1.3+(U{r}-0.6)/0.1*1.2,IF(U{r}<=0.8,2.5+(U{r}-0.7)/0.1*1,3.5))))/100', '0.00000'),
    ('psit', 'ψ(t∞)=2,5ψ1000', '', lambda r, i, c: f'=2.5*V{r}', '0.00000'),
    ('chi', 'χ = -ln(1-ψ)', '', lambda r, i, c: f'=-LN(1-W{r})', '0.00000'),
    ('eta', 'η = 1+e²Ac/Ic', '', lambda r, i, c: f'=1+D{r}^2*{Dd("Ac")}/{Dd("Ic")}', '0.0000'),
    ('dprog', 'Δσp progressivas', 'MPa', lambda r, i, c: f'=(ABS({Dd("ecs")})*{Dd("Ep")}+{Dd("ap28")}*ABS(T{r})*{Dd("phi")}+R{r}*X{r})/((1+X{r})+(1+0.5*{Dd("phi")})*{Dd("ap28")}*Y{r}*{Dd("rho_p")})', '0.0'),
    ('sinf', 'σp,∞', 'MPa', lambda r, i, c: f'=R{r}-Z{r}', '0.0'),
    ('Pinf', 'P∞', 'kN', lambda r, i, c: f'={Dd("Ap")}*AA{r}*1000', '0.0'),
    ('gpi', 'Pi/Pj', '', lambda r, i, c: f'=S{r}/{Dd("Pj")}', '0.0000'),
    ('gpp', 'P∞/Pi', '', lambda r, i, c: f'=AB{r}/S{r}', '0.0000'),
]
wsT, cTm = station_sheet('Tracado_Perdas', 'Traçado do cabo resultante e perdas de protensão (mecanismos físicos)', cT,
                         widths={'sL': 14, 'dk': 14, 'ss': 12, 'scp': 12},
                         note='Cabo resultante parabólico composto (vão + parábola reversa no apoio). Protensão simultânea pelas duas extremidades; atrito (NBR 6118 9.6.3.3.2), cravação, encurtamento elástico e perdas progressivas (NBR 6118 9.6.3.4.2).')

# ------------------------------------------------------------------ Hiperestatico
TPn = 'Tracado_Perdas'
cH = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('m', 'm(x) = momento da carga unitária em B', 'm', lambda r, i, c: f'=IF(A{r}<={L_},A{r}/2,(2*{L_}-A{r})/2)', '0.000'),
    ('Mi', 'M_iso,i = -Pi·cosα·e', 'kN·m', lambda r, i, c: f'=-{TPn}!S{r}*{TPn}!G{r}*{TPn}!D{r}', '0.00'),
    ('Mim', 'M_iso,i·m', '', lambda r, i, c: f'=C{r}*B{r}', '0.000'),
    ('Mf', 'M_iso,∞ = -P∞·cosα·e', 'kN·m', lambda r, i, c: f'=-{TPn}!AB{r}*{TPn}!G{r}*{TPn}!D{r}', '0.00'),
    ('Mfm', 'M_iso,∞·m', '', lambda r, i, c: f'=E{r}*B{r}', '0.000'),
    ('Hi', 'M_hip,i', 'kN·m', lambda r, i, c: f'=-$H$3*B{r}', '0.00'),
    ('Hf', 'M_hip,∞', 'kN·m', lambda r, i, c: f'=-$H$4*B{r}', '0.00'),
]
wsH, cHm = station_sheet('Hiperestatico', 'Efeito hiperestático da protensão (momento secundário) - viga de 2 vãos', cH,
                         widths={'m': 16},
                         note='Método das forças: retira-se o apoio B; R_B = 6·∫M_iso·m dx / L³ (EI constante). M_hip(x) = -R_B·m(x). Convenção: M positivo = tração embaixo.')
wsH['F3'] = 'R_B (ato) [kN]'; wsH['G3'] = 'ΣR'
wsH['F3'].font = fbb; wsH['F4'] = 'R_B (∞) [kN]'; wsH['F4'].font = fbb
wsH['H3'] = f'=6*{DX}*(SUM(D{FIRST}:D{LAST})-0.5*(D{FIRST}+D{LAST}))/{L_}^3'
wsH['H4'] = f'=6*{DX}*(SUM(F{FIRST}:F{LAST})-0.5*(F{FIRST}+F{LAST}))/{L_}^3'
wsH['G3'] = 'R_B>0: reação p/ cima'
wsH['H3'].fill = fill_res; wsH['H4'].fill = fill_res
wsH['H3'].number_format = '0.00'; wsH['H4'].number_format = '0.00'
wsH['I3'] = 'M_hip em B (ato) ='; wsH['J3'] = f'=-H3*{L_}/2'
wsH['I4'] = 'M_hip em B (∞) ='; wsH['J4'] = f'=-H4*{L_}/2'
for a in ('J3', 'J4'):
    wsH[a].fill = fill_res; wsH[a].number_format = '0.00'

# ------------------------------------------------------------------ Tensoes
W_ = Dd('W'); Ac_ = Dd('Ac')
cS = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('e', 'e', 'm', lambda r, i, c: f'={TPn}!D{r}', '0.000'),
    ('Ni', 'N_i = Pi·cosα', 'kN', lambda r, i, c: f'={TPn}!S{r}*{TPn}!G{r}', '0.0'),
    ('Nf', 'N_∞ = P∞·cosα', 'kN', lambda r, i, c: f'={TPn}!AB{r}*{TPn}!G{r}', '0.0'),
    ('MPi', 'M_P,i = M_iso+M_hip', 'kN·m', lambda r, i, c: f'=Hiperestatico!C{r}+Hiperestatico!G{r}', '0.0'),
    ('MPf', 'M_P,∞', 'kN·m', lambda r, i, c: f'=Hiperestatico!E{r}+Hiperestatico!H{r}', '0.0'),
    ('s1pi', 'σ1p (ato)', 'MPa', lambda r, i, c: f'=-C{r}/{Ac_}/1000+E{r}/{W_}/1000', '0.00'),
    ('s2pi', 'σ2p (ato)', 'MPa', lambda r, i, c: f'=-C{r}/{Ac_}/1000-E{r}/{W_}/1000', '0.00'),
    ('s1pf', 'σ1p (∞)', 'MPa', lambda r, i, c: f'=-D{r}/{Ac_}/1000+F{r}/{W_}/1000', '0.00'),
    ('s2pf', 'σ2p (∞)', 'MPa', lambda r, i, c: f'=-D{r}/{Ac_}/1000-F{r}/{W_}/1000', '0.00'),
    ('a1', 'σ1 ATO', 'MPa', lambda r, i, c: f'=G{r}+Esforcos!D{r}/{W_}/1000', '0.00'),
    ('a2', 'σ2 ATO', 'MPa', lambda r, i, c: f'=H{r}-Esforcos!D{r}/{W_}/1000', '0.00'),
    ('f1', 'σ1 CFQ (ELS-D)', 'MPa', lambda r, i, c: f'=I{r}+Esforcos!H{r}/{W_}/1000', '0.00'),
    ('f2', 'σ2 CFQ (ELS-D)', 'MPa', lambda r, i, c: f'=J{r}-Esforcos!I{r}/{W_}/1000', '0.00'),
    ('r1', 'σ1 CR (ELS-F)', 'MPa', lambda r, i, c: f'=I{r}+Esforcos!J{r}/{W_}/1000', '0.00'),
    ('r2', 'σ2 CR (ELS-F)', 'MPa', lambda r, i, c: f'=J{r}-Esforcos!K{r}/{W_}/1000', '0.00'),
    ('c2', 'σ2 CR (ELS-CE, M+)', 'MPa', lambda r, i, c: f'=J{r}-Esforcos!J{r}/{W_}/1000', '0.00'),
    ('c1', 'σ1 CR (ELS-CE, M-)', 'MPa', lambda r, i, c: f'=I{r}+Esforcos!K{r}/{W_}/1000', '0.00'),
    ('oka', 'Ato', '', lambda r, i, c: f'=IF(AND(K{r}<={Dd("ft_ato")},L{r}<={Dd("ft_ato")},K{r}>=-{Dd("fc_ato")},L{r}>=-{Dd("fc_ato")}),"OK","ARMAR")', None),
    ('okd', 'ELS-D', '', lambda r, i, c: f'=IF(AND(M{r}<=0,N{r}<=0),"OK","NÃO")', None),
    ('okf', 'ELS-F', '', lambda r, i, c: f'=IF(AND(O{r}<={Dd("ft_els")},P{r}<={Dd("ft_els")}),"OK","NÃO")', None),
    ('okc', 'ELS-CE', '', lambda r, i, c: f'=IF(AND(Q{r}>=-{Dd("fc_els")},R{r}>=-{Dd("fc_els")},K{r}>=-{Dd("fc_ato")}),"OK","NÃO")', None),
    ('Aato', 'As,ato (σs=250 MPa)', 'cm²', lambda r, i, c: (f'=MAX(IF(K{r}>{Dd("ft_ato")},0.5*K{r}*1000*{Dd("bw")}*({Dd("h")}*K{r}/(K{r}-L{r}))/250000*10000,0),IF(L{r}>{Dd("ft_ato")},0.5*L{r}*1000*{Dd("bw")}*({Dd("h")}*L{r}/(L{r}-K{r}))/250000*10000,0))'), '0.00'),
]
wsS, cSm = station_sheet('Tensoes', 'Verificação de tensões normais - ato, ELS-D (CFQ), ELS-F e ELS-CE (CR)', cS,
                         widths={'oka': 8, 'okd': 8, 'okf': 8, 'okc': 8},
                         note='σ1 = fibra inferior; σ2 = fibra superior. Tração positiva (+), compressão negativa (-). Hiperestático incluído nas forças de protensão.')
# resumo de extremos
wsS['A3'] = 'Extremos'; wsS['A3'].font = fbb
labels = [('K', 'ATO σ1 máx (tração)', 'MAX'), ('L', 'ATO σ2 máx (tração)', 'MAX'), ('M', 'CFQ σ1 máx', 'MAX'), ('N', 'CFQ σ2 máx', 'MAX'),
          ('O', 'CR σ1 máx', 'MAX'), ('P', 'CR σ2 máx', 'MAX'), ('Q', 'CR σ2 mín (compr.)', 'MIN'), ('R', 'CR σ1 mín (compr.)', 'MIN')]
for col, lab, fn in labels:
    wsS[f'{col}4'] = f'={fn}({col}{FIRST}:{col}{LAST})'; wsS[f'{col}4'].fill = fill_res; wsS[f'{col}4'].number_format = '0.00'
    wsS[f'{col}3'] = fn; wsS[f'{col}3'].font = Font(size=8)
wsS['K5'] = f'=MIN(K{FIRST}:K{LAST},L{FIRST}:L{LAST})'; wsS['K5'].number_format = '0.00'; wsS['K5'].fill = fill_res
wsS['J5'] = 'ATO compr. mín'; wsS['J5'].font = Font(size=8)

# ------------------------------------------------------------------ Fuso
Wd = W_; kd = Dd('k')
Mg0c = lambda r: f'Esforcos!D{r}'
def f_ft_ato(): return f'{Dd("ft_ato")}*1000'
cF = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('e', 'e (cabo)', 'm', lambda r, i, c: f'={TPn}!D{r}', '0.000'),
    ('elp', 'e_lp,∞ = e - M_hip/N', 'm', lambda r, i, c: f'=B{r}-Hiperestatico!H{r}/Tensoes!D{r}', '0.000'),
    ('L1', 'inf: ato tração σ1', 'm', lambda r, i, c: f'=({Mg0c(r)}-{Dd("ft_ato")}*1000*{Wd})/Tensoes!C{r}-{kd}', '0.000'),
    ('L2', 'inf: ato compr. σ2', 'm', lambda r, i, c: f'=({Mg0c(r)}-{Dd("fc_ato")}*1000*{Wd})/Tensoes!C{r}+{kd}', '0.000'),
    ('L3', 'inf: ELS-D σ1 (CFQ)', 'm', lambda r, i, c: f'=Esforcos!H{r}/Tensoes!D{r}-{kd}', '0.000'),
    ('L4', 'inf: ELS-F σ1 (CR)', 'm', lambda r, i, c: f'=(Esforcos!J{r}-{Dd("ft_els")}*1000*{Wd})/Tensoes!D{r}-{kd}', '0.000'),
    ('L5', 'inf: ELS-CE σ2 (CR)', 'm', lambda r, i, c: f'=(Esforcos!J{r}-{Dd("fc_els")}*1000*{Wd})/Tensoes!D{r}+{kd}', '0.000'),
    ('U1', 'sup: ato compr. σ1', 'm', lambda r, i, c: f'=({Mg0c(r)}+{Dd("fc_ato")}*1000*{Wd})/Tensoes!C{r}-{kd}', '0.000'),
    ('U2', 'sup: ato tração σ2', 'm', lambda r, i, c: f'=({Mg0c(r)}+{Dd("ft_ato")}*1000*{Wd})/Tensoes!C{r}+{kd}', '0.000'),
    ('U3', 'sup: ELS-D σ2 (CFQ)', 'm', lambda r, i, c: f'=Esforcos!I{r}/Tensoes!D{r}+{kd}', '0.000'),
    ('U4', 'sup: ELS-F σ2 (CR)', 'm', lambda r, i, c: f'=({Dd("ft_els")}*1000*{Wd}+Esforcos!K{r})/Tensoes!D{r}+{kd}', '0.000'),
    ('U5', 'sup: ELS-CE σ1 (CR)', 'm', lambda r, i, c: f'=(Esforcos!K{r}+{Dd("fc_els")}*1000*{Wd})/Tensoes!D{r}-{kd}', '0.000'),
    ('lo', 'Limite inferior do fuso (e mín)', 'm', lambda r, i, c: f'=MAX(D{r}:H{r})', '0.000'),
    ('up', 'Limite superior do fuso (e máx)', 'm', lambda r, i, c: f'=MIN(I{r}:M{r})', '0.000'),
    ('ok', 'e_lp no fuso?', '', lambda r, i, c: f'=IF(AND(C{r}>=N{r},C{r}<=O{r}),"SIM",IF(AND(C{r}>=MAX(F{r}:H{r}),C{r}<=MIN(K{r}:M{r})),"SIM c/ As,ato","NÃO"))', None),
    ('cov', 'Limite de cobrimento', 'm', lambda r, i, c: f'={Dd("emax")}', '0.000'),
]
wsF, cFm = station_sheet('Fuso', 'Fuso limite - região admissível da linha de pressões da protensão (e positivo abaixo do CG)', cF,
                         widths={'ok': 14},
                         note='Limites de e_lp (m): ato (Pi, Mg0) e serviço (P∞, CFQ/CR). A linha de pressões e_lp = e - M_hip/N deve permanecer entre os limites. Quando o limite de tração no ato é violado, adota-se armadura passiva (As,ato).')

# ------------------------------------------------------------------ Cortante
Tn = 'Tensoes'
cC = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('Vpi', 'Vp,iso = -P∞·senα', 'kN', lambda r, i, c: f'=-{TPn}!AB{r}*{TPn}!H{r}', '0.0'),
    ('Vh', 'V_hip', 'kN', lambda r, i, c: f'=IF(A{r}<={L_},-Hiperestatico!$H$4/2,Hiperestatico!$H$4/2)', '0.00'),
    ('Vp', 'Vp total', 'kN', lambda r, i, c: f'=B{r}+C{r}', '0.0'),
    ('Vg', 'Vg', 'kN', lambda r, i, c: f'=Esforcos!L{r}', '0.0'),
    ('Vqm', 'Vq+', 'kN', lambda r, i, c: f'=Esforcos!M{r}', '0.0'),
    ('Vqn', 'Vq-', 'kN', lambda r, i, c: f'=Esforcos!N{r}', '0.0'),
    ('c1', 'C1 0,9Vp+1,4Vg', 'kN', lambda r, i, c: f'=0.9*D{r}+1.4*E{r}', '0.0'),
    ('c2', 'C2 +1,4Vq+', 'kN', lambda r, i, c: f'=0.9*D{r}+1.4*E{r}+1.4*F{r}', '0.0'),
    ('c3', 'C2b +1,4Vq-', 'kN', lambda r, i, c: f'=0.9*D{r}+1.4*E{r}+1.4*G{r}', '0.0'),
    ('c4', 'C3 1,2Vp+1,0Vg', 'kN', lambda r, i, c: f'=1.2*D{r}+E{r}', '0.0'),
    ('c5', 'C4 +1,4Vq+', 'kN', lambda r, i, c: f'=1.2*D{r}+E{r}+1.4*F{r}', '0.0'),
    ('c6', 'C4b +1,4Vq-', 'kN', lambda r, i, c: f'=1.2*D{r}+E{r}+1.4*G{r}', '0.0'),
    ('Vsd', 'Vsd = máx |C|', 'kN', lambda r, i, c: f'=MAX(ABS(H{r}),ABS(I{r}),ABS(J{r}),ABS(K{r}),ABS(L{r}),ABS(M{r}))', '0.0'),
    ('Msag', 'Msd+ = 1,4(Mg+Mq+)', 'kN·m', lambda r, i, c: f'=1.4*Esforcos!J{r}', '0.0'),
    ('Mhog', 'Msd- = 1,4(Mg+Mq-)', 'kN·m', lambda r, i, c: f'=1.4*Esforcos!K{r}', '0.0'),
    ('Msd0', '|Msd|', 'kN·m', lambda r, i, c: f'=MAX(ABS(O{r}),ABS(P{r}))', '0.0'),
    ('Msdadj', '|Msd| (máx. entre seções adjacentes)', 'kN·m', lambda r, i, c: (f'=MAX(Q{r},Q{r+1})' if i == 0 else (f'=MAX(Q{r},Q{r-1})' if i == NST - 1 else f'=MAX(Q{r},Q{r-1},Q{r+1})')), '0.0'),
    ('M0', 'M0 (descompressão, γp=0,9)', 'kN·m', lambda r, i, c: f'=IF(ABS(O{r})>=ABS(P{r}),0.9*({Tn}!D{r}*{kd}-{Tn}!F{r}),0.9*({Tn}!D{r}*{kd}+{Tn}!F{r}))', '0.0'),
    ('beta', 'β = 1+M0/Msd ≤ 2', '', lambda r, i, c: f'=MIN(1+S{r}/MAX(R{r},0.000001),2)', '0.000'),
    ('tcc', 'τc,cor = β·τc', 'MPa', lambda r, i, c: f'=T{r}*{Dd("tauc")}', '0.000'),
    ('twd', 'τwd = Vsd/(bw,red·d)', 'MPa', lambda r, i, c: f'=N{r}/({Dd("bwred")}*{Dd("dd")})/1000', '0.000'),
    ('td', 'τd = (τwd-τc,cor)/0,9', 'MPa', lambda r, i, c: f'=(V{r}-U{r})/0.9', '0.000'),
    ('asw', 'asw/n (por ramo)', 'cm²/m', lambda r, i, c: f'=MAX(W{r},0)/({Dd("fyk")}/{Dd("gs")})*{Dd("bwred")}*10000/2', '0.00'),
    ('aswmin', 'asw,mín/n', 'cm²/m', lambda r, i, c: f'={Dd("rhomin")}*{Dd("bw")}*10000/2', '0.00'),
    ('aswc', 'asw,calc/n', 'cm²/m', lambda r, i, c: f'=MAX(X{r},Y{r})', '0.00'),
    ('Vu', 'VRd2 = τu·bw,red·d', 'kN', lambda r, i, c: f'={Dd("tauu")}*1000*{Dd("bwred")}*{Dd("dd")}', '0.0'),
    ('ok', 'Vsd ≤ VRd2?', '', lambda r, i, c: f'=IF(N{r}<=AA{r},"OK","NÃO")', None),
]
wsC, cCm = station_sheet('Cortante', 'Força cortante - protensão (Vp), combinações NBR 6118 e estribos (Modelo I)', cC,
                         widths={'Msdadj': 14},
                         note='Metodologia da aula "Força cortante" (Vp = -P·senα; Vsd = 0,9Vp+1,4Vg(+1,4Vq) ou 1,2Vp+1,0Vg(+1,4Vq); β = 1+M0/Msd; τc,cor = β·τc). Estribos por ramo (2 ramos).')

# ------------------------------------------------------------------ ELU
wsU = wb.create_sheet('ELU')
wsU['A1'] = 'Verificação à ruptura (ELU - flexão)'; wsU['A1'].font = Font(bold=True, size=13, color=NAVY)
wsU['A2'] = 'Pré-alongamento (neutralização), diagrama bilinear do aço de protensão (patamar em fpyd), retângulo de tensões 0,8x·0,85·fcd. A' + chr(39) + ' = armadura passiva comprimida (entrada).'
wsU['A2'].font = Font(italic=True, size=9)
hdr = ['Grandeza', 'Apoio B (x=10 m)', 'Vão (x=4,0 m)', 'Unid.']
for j, hh in enumerate(hdr, 1):
    c = wsU.cell(4, j, hh); c.font = fh; c.fill = fill_h
rB = FIRST + 20; rV = FIRST + 8
items = [
    ('Linha de estação', lambda r, col: r, '0'),
    ('Face comprimida', lambda r, col: 'inferior' if col == 'B' else 'superior', None),
    ('e (cabo, + abaixo CG)', lambda r, col: f'={TPn}!D{r}', '0.000'),
    ('dp (cabo à face comprimida)', lambda r, col: (f'={Dd("h")}/2-{col}7' if col == 'B' else f'={Dd("h")}/2+{col}7'), '0.000'),
    ('P∞ no cabo', lambda r, col: f'={TPn}!AB{r}', '0.0'),
    ('σp,∞', lambda r, col: f'={TPn}!AA{r}', '0.0'),
    ('εp0 = σ∞/Ep + P∞/(Ac·Ecs)·(1+e²Ac/Ic)', lambda r, col: f'={col}10/{Dd("Ep")}+{col}9/({Dd("Ac")}*{Dd("Ecs")}*1000)*(1+{col}7^2*{Dd("Ac")}/{Dd("Ic")})', '0.00000'),
    ("A' (compressão, passiva)", lambda r, col: (4.91 if col == 'B' else 2.45), '0.00'),
    ("d' (A' à face comprimida)", lambda r, col: 0.05, '0.000'),
    ('Tp = Ap·fpyd', lambda r, col: f'={Dd("Ap")}*{Dd("fpyd")}*1000', '0.0'),
    ("Cs = A'·fyd", lambda r, col: f'={col}12*0.0001*{Dd("fyk")}/{Dd("gs")}*1000', '0.0'),
    ('x = (Tp - Cs)/(0,85·fcd·0,8·bw)', lambda r, col: f'=({col}14-{col}15)/(0.85*{Dd("fcd")}*1000*0.8*{Dd("bw")})', '0.0000'),
    ('x/dp', lambda r, col: f'={col}16/{col}8', '0.000'),
    ('εp = εp0 + 3,5‰·(dp-x)/x', lambda r, col: f'={col}11+0.0035*({col}8-{col}16)/{col}16', '0.00000'),
    ('εpy = fpyd/Ep  (εp ≥ εpy ?)', lambda r, col: f'=IF({col}18>={Dd("fpyd")}/{Dd("Ep")},"escoamento OK","VERIFICAR")', None),
    ("ε's = 3,5‰·(x-d')/x ≥ εyd ?", lambda r, col: f'=IF(0.0035*({col}16-{col}13)/{col}16>={Dd("fyk")}/{Dd("gs")}/{Dd("Es")},"escoa","VERIFICAR")', None),
    ('Cc = Tp - Cs', lambda r, col: f'={col}14-{col}15', '0.0'),
    ("MRd = Cc(dp-0,4x)+Cs(dp-d')", lambda r, col: f'={col}21*({col}8-0.4*{col}16)+{col}15*({col}8-{col}13)', '0.0'),
    ('M_hip,∞ na seção', lambda r, col: f'=Hiperestatico!H{r}', '0.0'),
    ('Msd externo = 1,4(Mg+Mq)', lambda r, col: (f'=1.4*Esforcos!K{r}' if col == 'B' else f'=1.4*Esforcos!J{r}'), '0.0'),
    ('γp (1,2 desfavorável / 0,9 favorável)', lambda r, col: (f'=IF({col}23>0,0.9,1.2)' if col == 'B' else f'=IF({col}23<0,0.9,1.2)'), '0.0'),
    ('Msd = externo + γp·M_hip', lambda r, col: f'={col}24+{col}25*{col}23', '0.0'),
    ('|Msd| ≤ MRd ?', lambda r, col: f'=IF(ABS({col}26)<={col}22,"OK","NÃO")', None),
    ('MRd/|Msd|', lambda r, col: f'={col}22/ABS({col}26)', '0.00'),
    ("Ductilidade x/d ≤ 0,45 ?", lambda r, col: f'=IF({col}17<=0.45,"OK","VERIFICAR A\'")', None),
    ('Md,mín = 0,8·W·fctk,sup (NBR 17.3.5.2.1)', lambda r, col: f'=0.8*{Dd("W")}*1.3*{Dd("fctm")}*1000', '0.0'),
]
for n, (lab, fn, fmt) in enumerate(items):
    row = 5 + n
    wsU.cell(row, 1, lab).font = fb
    for col, rr_ in (('B', rB), ('C', rV)):
        v = fn(rr_, col)
        c = wsU[f'{col}{row}']
        c.value = v; c.font = fb; c.border = box
        if fmt: c.number_format = fmt
        if lab.startswith("A'") or lab.startswith("d'"):
            c.fill = fill_in
        if lab.startswith('MRd') or lab.startswith('Msd =') or lab.startswith('|Msd|'):
            c.fill = fill_res
wsU.column_dimensions['A'].width = 48; wsU.column_dimensions['B'].width = 18; wsU.column_dimensions['C'].width = 18
# rows in `items` are 5.. ; the formulas above reference fixed rows (7..): build mapping check
# (items index n -> row 5+n): row7=e, row8=dp, row9=P, row10=sig, row11=eps0, row12=A', row13=d', row14=Tp, row15=Cs, row16=x,
# row17=x/dp, row18=eps_p, row19=check, row20=check2->Cc? (see below)

# ------------------------------------------------------------------ Alongamento
wsA = wb.create_sheet('Alongamento')
wsA['A1'] = 'Alongamento final dos cabos'; wsA['A1'].font = Font(bold=True, size=13, color=NAVY)
wsA['A2'] = 'Protensão simultânea pelas duas extremidades até σp,j. Δl por extremidade = ∫σp(s)/Ep ds ao longo do cabo (x de 0 a L), com σp(s) do atrito antes da cravação.'
wsA['A2'].font = Font(italic=True, size=9)
sumrows = [
    ('Δl por extremidade (teórico, pré-cravação)', f'=SUMPRODUCT(({TPn}!M{FIRST+1}:M{FIRST+20}+{TPn}!M{FIRST}:M{FIRST+19})/2,{TPn}!K{FIRST+1}:K{FIRST+20})/{Dd("Ep")}*1000', 'mm'),
    ('Δl total do cabo (2 extremidades)', '=2*B4', 'mm'),
    ('Tolerância ±5 %  (mín.)', '=0.95*B4', 'mm'),
    ('Tolerância ±5 %  (máx.)', '=1.05*B4', 'mm'),
    ('Cravação das cunhas (por extremidade)', f'={Dd("slip")}*1000', 'mm'),
    ('Alongamento líquido por extremidade após cravação', '=B4-B8', 'mm'),
    ('Comprimento do cabo (arco, total)', f'={TPn}!L{LAST}', 'm'),
    ('Ângulo total desviado Σα', f'={TPn}!J{LAST}', 'rad'),
    ('Perda por atrito até o meio do cabo (apoio B)', f'=1-{TPn}!M{FIRST+20}/{Dd("sj")}', ''),
    ('Alcance da cravação xa', f'={Dd("xa")}', 'm'),
]
wsA['A3'] = 'Grandeza'; wsA['B3'] = 'Valor'; wsA['C3'] = 'Unid.'
for cc in ('A3', 'B3', 'C3'):
    wsA[cc].font = fh; wsA[cc].fill = fill_h
for n, (lab, f, u) in enumerate(sumrows):
    wsA.cell(4 + n, 1, lab).font = fb
    c = wsA.cell(4 + n, 2, f); c.fill = fill_res; c.number_format = '0.00'; c.border = box
    wsA.cell(4 + n, 3, u).font = fb
wsA.column_dimensions['A'].width = 56; wsA.column_dimensions['B'].width = 14

# ------------------------------------------------------------------ Flecha PC
EI = f'({Dd("Ecs")}*1000*{Dd("Ic")})'
def block(name_prefix, mcol_fn):
    pass
cP = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('Mg', 'Mg', 'kN·m', lambda r, i, c: f'=Esforcos!E{r}', '0.0'),
    ('Mq2', 'ψ2·Mq (ambos vãos)', 'kN·m', lambda r, i, c: f'={Dd("psi2")}*{Mboth(q_, "Esforcos!C"+str(r))}', '0.0'),
    ('MP', 'M_P,∞', 'kN·m', lambda r, i, c: f'=Tensoes!F{r}', '0.0'),
    ('Mqp', 'M_qp líquido = Mg+ψ2Mq+M_P', 'kN·m', lambda r, i, c: f'=B{r}+C{r}+D{r}', '0.0'),
    ('k1', 'κ_g', '1/m', lambda r, i, c: f'=-B{r}/{EI}', '0.0000000'),
    ('k2', 'κ_P', '1/m', lambda r, i, c: f'=-D{r}/{EI}', '0.0000000'),
    ('k3', 'κ_qp', '1/m', lambda r, i, c: f'=-E{r}/{EI}', '0.0000000'),
]
# integracao: para cada kappa, colunas slope (S) e y (Y) e y corrigido (mm)
base = len(cP)
def add_int(cols, kcol, tag):
    sl = f'sl_{tag}'; yy = f'y_{tag}'; yc = f'yc_{tag}'
    n0 = len(cols)
    # letras: calculado dinamicamente via colmap, entao usa funcoes que leem colmap
    cols.append((sl, f'θ_{tag}', 'rad', (lambda kc: (lambda r, i, c: 0 if i == 0 else f'={c[sl]}{r-1}+0.5*({c[kc]}{r}+{c[kc]}{r-1})*({c["x"]}{r}-{c["x"]}{r-1})'))(kcol), '0.000000'))
    cols.append((yy, f'y_{tag} (bruto)', 'm', (lambda s_: (lambda r, i, c: 0 if i == 0 else f'={c[yy]}{r-1}+0.5*({c[s_]}{r}+{c[s_]}{r-1})*({c["x"]}{r}-{c["x"]}{r-1})'))(sl), '0.0000000'))
    cols.append((yc, f'δ_{tag} (+ para baixo)', 'mm', (lambda y_: (lambda r, i, c: f'=({c[y_]}{r}-{c[y_]}${FIRST+20}*{c["x"]}{r}/{L_})*1000'))(yy), '0.000'))
add_int(cP, 'k1', 'g'); add_int(cP, 'k2', 'P'); add_int(cP, 'k3', 'qp')
cP.append(('dinf', 'δ_qp,∞ = δ_qp·(1+φ)', 'mm', lambda r, i, c: f'={c["yc_qp"]}{r}*(1+{Dd("phi")})', '0.000'))
wsP, cPm = station_sheet('Flecha_PC', 'ELS-DEF - flecha da viga protendida (seção não fissurada, Ecs·Ic)', cP,
                         widths={'Mqp': 16, 'yc_g': 12, 'yc_P': 12, 'yc_qp': 12, 'dinf': 14},
                         note='y"=-M/EI integrado duas vezes; y(0)=y(L)=0. Combinação quase permanente: Mg + ψ2·Mq + efeito da protensão P∞. Deslocamento positivo = para baixo. Limite visual: L/250.')
wsP['M3'] = 'máx |δ_g|'; wsP['N3'] = f'=MAX(ABS(MIN({cPm["yc_g"]}{FIRST}:{cPm["yc_g"]}{FIRST+20})),MAX({cPm["yc_g"]}{FIRST}:{cPm["yc_g"]}{FIRST+20}))'
wsP['M4'] = 'máx |δ_P|'; wsP['N4'] = f'=MAX(ABS(MIN({cPm["yc_P"]}{FIRST}:{cPm["yc_P"]}{FIRST+20})),MAX({cPm["yc_P"]}{FIRST}:{cPm["yc_P"]}{FIRST+20}))'
wsP['M5'] = 'δ_qp (mín.)'; wsP['N5'] = f'=MIN({cPm["yc_qp"]}{FIRST}:{cPm["yc_qp"]}{FIRST+20})'
wsP['P3'] = 'δ_qp,∞ (mín.)'; wsP['Q3'] = f'=MIN({cPm["dinf"]}{FIRST}:{cPm["dinf"]}{FIRST+20})'
wsP['P4'] = 'Limite L/250 [mm]'; wsP['Q4'] = f'={L_}*1000/250'
wsP['P5'] = 'f1 (modo antissimétrico) [Hz]'; wsP['Q5'] = f'=PI()/(2*{L_}^2)*SQRT({Dd("Ecs")}*1000000*{Dd("Ic")}/({Dd("g")}/9.81*1000))'
for a in ('N3', 'N4', 'N5', 'Q3', 'Q4', 'Q5'):
    wsP[a].fill = fill_res; wsP[a].number_format = '0.00'

# ------------------------------------------------------------------ CA
wsR = wb.create_sheet('CA')
wsR['A1'] = 'Comparação com concreto armado - mesma seção 30×90, mesmas cargas'; wsR['A1'].font = Font(bold=True, size=13, color=NAVY)
wsR['A2'] = 'ELU: bloco retangular 0,8x·0,85·fcd (sem redistribuição). ELS-W: NBR 6118 17.3.3.2 (CF, ψ1). Deflexão: Branson com integração (aba CA_Flecha). Entradas em amarelo.'
wsR['A2'].font = Font(italic=True, size=9)
for j, hh in enumerate(['Grandeza', 'Apoio B (M-)', 'Vão (M+)', 'Unid.'], 1):
    c = wsR.cell(4, j, hh); c.font = fh; c.fill = fill_h
ca_items = [
    ('d (altura útil, 2 camadas)', lambda col: 0.82, '0.000', 'in'),
    ('Md = 1,4·(Mg+Mq)', lambda col: (f'=1.4*ABS(Esforcos!K{rB})' if col == 'B' else f'=1.4*Esforcos!J{rV}'), '0.0', ''),
    ('Kmd = Md/(bw·d²·fcd)', lambda col: f'={col}6/({Dd("bw")}*{col}5^2*{Dd("fcd")}*1000)', '0.0000', ''),
    ('x/d = 1,25·[1-√(1-2Kmd/0,85)]', lambda col: f'=1.25*(1-SQRT(1-2*{col}7/0.85))', '0.000', ''),
    ('x', lambda col: f'={col}8*{col}5', '0.000', ''),
    ('As,calc = Md/[fyd·(d-0,4x)]', lambda col: f'={col}6/({Dd("fyk")}/{Dd("gs")}*1000*({col}5-0.4*{col}9))*10000', '0.00', ''),
    ('Barras adotadas (nº de Ø20)', lambda col: (8 if col == 'B' else 6), '0', 'in'),
    ('As,adot', lambda col: f'={col}11*3.1416', '0.00', ''),
    ('Ductilidade x/d ≤ 0,45', lambda col: f'=IF({col}8<=0.45,"OK","NÃO")', None, ''),
    ('Momento de fissuração Mr = 1,5·fct,m·Ic/yt', lambda col: f'=1.5*{Dd("fctm")}*1000*{Dd("Ic")}/({Dd("h")}/2)', '0.0', ''),
    ('Ma (CF) = |Mg + ψ1·Mq|', lambda col: (f'=ABS(Esforcos!I{rB})' if col == 'B' else f'=Esforcos!H{rV}'), '0.0', ''),
    ('αe = Es/Ecs', lambda col: f'={Dd("Es")}/{Dd("Ecs")}', '0.000', ''),
    ('x_II (estádio II)', lambda col: f'=(-{col}16*{col}12/10000+SQRT(({col}16*{col}12/10000)^2+2*{Dd("bw")}*{col}16*{col}12/10000*{col}5))/{Dd("bw")}', '0.0000', ''),
    ('I_II', lambda col: f'={Dd("bw")}*{col}17^3/3+{col}16*{col}12/10000*({col}5-{col}17)^2', '0.000000', ''),
    ('σs = αe·Ma·(d-x_II)/I_II', lambda col: f'={col}16*{col}15*({col}5-{col}17)/{col}18/1000', '0.0', ''),
    ('ρri = As/Acri', lambda col: f'={col}12/10000/({Dd("bw")}*(0.055+7.5*0.02))', '0.0000', ''),
    ('wk1 = φ/(12,5η1)·σs/Es·3σs/fct,m', lambda col: f'=20/(12.5*2.25)*{col}19/{Dd("Es")}*3*{col}19/{Dd("fctm")}', '0.000', ''),
    ('wk2 = φ/(12,5η1)·σs/Es·(4/ρri+45)', lambda col: f'=20/(12.5*2.25)*{col}19/{Dd("Es")}*(4/{col}20+45)', '0.000', ''),
    ('wk = mín(wk1, wk2)  [mm]  (limite 0,3)', lambda col: f'=MIN({col}21,{col}22)', '0.000', ''),
    ('Vsd (CA) = 1,4(Vg+Vq)', lambda col: (f'=1.4*ABS(Esforcos!L{rB}+Esforcos!N{rB})' if col == 'B' else f'=1.4*(Esforcos!L{FIRST}+Esforcos!M{FIRST})'), '0.0', ''),
    ('Vc = 0,6·fctd·bw·d', lambda col: f'=0.6*{Dd("fctd")}*1000*{Dd("bw")}*{col}5', '0.0', ''),
    ('asw/s total (2 ramos) = (Vsd-Vc)/(0,9·d·fywd)', lambda col: f'=MAX({col}24-{col}25,0)/(0.9*{col}5*{Dd("fyk")}/{Dd("gs")}*1000)*10000', '0.00', ''),
    ('asw,mín total', lambda col: f'={Dd("rhomin")}*{Dd("bw")}*10000', '0.00', ''),
]
for n, (lab, fn, fmt, kind) in enumerate(ca_items):
    row = 5 + n
    wsR.cell(row, 1, lab).font = fb
    for col in ('B', 'C'):
        c = wsR[f'{col}{row}']; c.value = fn(col); c.font = fb; c.border = box
        if fmt: c.number_format = fmt
        c.fill = fill_in if kind == 'in' else (fill_res if lab.startswith(('As,', 'wk =')) else PatternFill())
wsR.column_dimensions['A'].width = 52; wsR.column_dimensions['B'].width = 16; wsR.column_dimensions['C'].width = 16
# Fixed-row consistency note: items index->row mapping validated in tests.

# ------------------------------------------------------------------ CA_Flecha
cK = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('Mqp', 'M_qp = Mg+ψ2Mq', 'kN·m', lambda r, i, c: f'=Esforcos!E{r}+{Dd("psi2")}*{Mboth(q_, "Esforcos!C"+str(r))}', '0.0'),
    ('As', 'As (cm²)', 'cm²', lambda r, i, c: f'=IF(B{r}>=0,CA!$C$12,CA!$B$12)', '0.00'),
    ('x2', 'x_II', 'm', lambda r, i, c: f'=(-CA!$B$16*C{r}/10000+SQRT((CA!$B$16*C{r}/10000)^2+2*{Dd("bw")}*CA!$B$16*C{r}/10000*CA!$B$5))/{Dd("bw")}', '0.0000'),
    ('I2', 'I_II', 'm⁴', lambda r, i, c: f'={Dd("bw")}*D{r}^3/3+CA!$B$16*C{r}/10000*(CA!$B$5-D{r})^2', '0.000000'),
    ('Ie', 'Ie (Branson)', 'm⁴', lambda r, i, c: f'=IF(ABS(B{r})<=CA!$B$14,{Dd("Ic")},MIN((CA!$B$14/ABS(B{r}))^3*{Dd("Ic")}+(1-(CA!$B$14/ABS(B{r}))^3)*E{r},{Dd("Ic")}))', '0.000000'),
    ('kap', 'κ = -M/(Ecs·Ie)', '1/m', lambda r, i, c: f'=-B{r}/({Dd("Ecs")}*1000*F{r})', '0.0000000'),
]
def add_int2(cols):
    cols.append(('sl', 'θ', 'rad', lambda r, i, c: 0 if i == 0 else f'={c["sl"]}{r-1}+0.5*({c["kap"]}{r}+{c["kap"]}{r-1})*({c["x"]}{r}-{c["x"]}{r-1})', '0.000000'))
    cols.append(('yy', 'y bruto', 'm', lambda r, i, c: 0 if i == 0 else f'={c["yy"]}{r-1}+0.5*({c["sl"]}{r}+{c["sl"]}{r-1})*({c["x"]}{r}-{c["x"]}{r-1})', '0.0000000'))
    cols.append(('yc', 'δ inst. (mm)', 'mm', lambda r, i, c: f'=({c["yy"]}{r}-{c["yy"]}${FIRST+20}*{c["x"]}{r}/{L_})*1000', '0.000'))
add_int2(cK)
cK.append(('yt', 'δ total = δ·(1+αf) (mm)', 'mm', lambda r, i, c: f'={c["yc"]}{r}*(1+$K$3)', '0.000'))
wsK, cKm = station_sheet('CA_Flecha', 'Concreto armado - flecha (Branson, rigidez variável)', cK,
                         widths={'Ie': 12, 'yt': 14},
                         note='αf = Δξ/(1+50ρ\'), Δξ = 2 - 0,68·0,996 (t0 = 1 mês), ρ\' = 0 (conservador).')
wsK['J3'] = 'αf ='; wsK['K3'] = '=(2-0.68*0.996)/(1+50*0)'; wsK['K3'].fill = fill_res; wsK['K3'].number_format = '0.000'
wsK['J4'] = 'δ inst. máx (mm)'; wsK['K4'] = f'=MAX({cKm["yc"]}{FIRST}:{cKm["yc"]}{FIRST+20})'
wsK['J5'] = 'δ total máx (mm)'; wsK['K5'] = f'=MAX({cKm["yt"]}{FIRST}:{cKm["yt"]}{FIRST+20})'
for a in ('K4', 'K5'):
    wsK[a].fill = fill_res; wsK[a].number_format = '0.00'


# ------------------------------------------------------------------ Niveis de protensao (NBR 6118, Tab. 13.4)
cN = [
    ('x', 'x', 'm', lambda r, i, c: f'={TPn}!B{r}', '0.0'),
    ('q1', 'σ1 CQP (Mg+ψ2Mq+)', 'MPa', lambda r, i, c: f'=Tensoes!I{r}+(Esforcos!E{r}+{Dd("psi2")}*Esforcos!F{r})/{W_}/1000', '0.00'),
    ('q2', 'σ2 CQP (Mg+ψ2Mq-)', 'MPa', lambda r, i, c: f'=Tensoes!J{r}-(Esforcos!E{r}+{Dd("psi2")}*Esforcos!G{r})/{W_}/1000', '0.00'),
    ('f1', 'σ1 CF (ψ1)', 'MPa', lambda r, i, c: f'=Tensoes!M{r}', '0.00'),
    ('f2', 'σ2 CF (ψ1)', 'MPa', lambda r, i, c: f'=Tensoes!N{r}', '0.00'),
    ('r1', 'σ1 CR', 'MPa', lambda r, i, c: f'=Tensoes!O{r}', '0.00'),
    ('r2', 'σ2 CR', 'MPa', lambda r, i, c: f'=Tensoes!P{r}', '0.00'),
    ('n3d', 'N3 ELS-D (CF ≤ 0)', '', lambda r, i, c: f'=IF(AND(D{r}<=0,E{r}<=0),"OK","NÃO")', None),
    ('n3f', 'N3 ELS-F (CR ≤ ft)', '', lambda r, i, c: f'=IF(AND(F{r}<={Dd("ft_els")},G{r}<={Dd("ft_els")}),"OK","NÃO")', None),
    ('n2d', 'N2 ELS-D (CQP ≤ 0)', '', lambda r, i, c: f'=IF(AND(B{r}<=0,C{r}<=0),"OK","NÃO")', None),
    ('n2f', 'N2 ELS-F (CF ≤ ft)', '', lambda r, i, c: f'=IF(AND(D{r}<={Dd("ft_els")},E{r}<={Dd("ft_els")}),"OK","NÃO")', None),
    ('n1', 'N1 seção sob CF', '', lambda r, i, c: f'=IF(MAX(D{r},E{r})>{Dd("ft_els")},"fissurada: verificar wk","não fissurada")', None),
]
wsN, cNm = station_sheet('Niveis', 'Níveis de protensão - verificação do projeto atual (NBR 6118, Tab. 13.4, pós-tração, CAA II)', cN,
                         widths={'n1': 24, 'n3d': 11, 'n3f': 11, 'n2d': 11, 'n2f': 11},
                         note='Nível 3 (completa): ELS-D na comb. frequente e ELS-F na rara. Nível 2 (limitada): ELS-D na quase permanente e ELS-F na frequente. Nível 1 (parcial): só ELS-W (wk ≤ 0,2 mm, CAA II) na frequente. σ1 = fibra inferior; σ2 = superior; tração +.')
wsN['A3'] = 'Atende?'; wsN['A3'].font = fbb
for col, lab in (('H', 'N3'), ('I', 'N3'), ('J', 'N2'), ('K', 'N2')):
    wsN[f'{col}3'] = f'=IF(COUNTIF({col}{FIRST}:{col}{LAST},"NÃO")=0,"ATENDE","NÃO ATENDE")'
    wsN[f'{col}3'].fill = fill_res; wsN[f'{col}3'].font = fb
wsN['L3'] = f'=IF(COUNTIF(L{FIRST}:L{LAST},"fissurada*")=0,"sem fissuras","há trechos fissurados")'
wsN['L3'].fill = fill_res; wsN['L3'].font = fb
wsN['A4'] = 'σ tração máx.'; wsN['A4'].font = fbb
for col in ('B', 'C', 'D', 'E', 'F', 'G'):
    wsN[f'{col}4'] = f'=MAX({col}{FIRST}:{col}{LAST})'; wsN[f'{col}4'].fill = fill_res; wsN[f'{col}4'].number_format = '0.00'

wsQ = wb.create_sheet('Niveis_Comparativo')
wsQ['A1'] = 'Comparativo dos níveis de protensão - mesma seção 30×90, C30, CAA II, traçado e = 28 cm'; wsQ['A1'].font = Font(bold=True, size=13, color=NAVY)
wsQ['A2'] = ('Cordoalhas por cabo e áreas são fórmulas; margens, P∞, armaduras e flechas vêm do motor de cálculo (Python), '
             'rodado com a mesma metodologia das demais abas. Para conferir um nível, altere ncord em Dados e leia a aba Niveis.')
wsQ['A2'].font = Font(italic=True, size=9)
heads = ['Grandeza', 'Unid.', 'Nível 3 - completa', 'Nível 2 - limitada', 'Nível 1 - parcial']
for j, h_ in enumerate(heads, start=1):
    c_ = wsQ.cell(4, j, h_); c_.font = fh; c_.fill = fill_h; c_.alignment = Alignment(horizontal='center', wrap_text=True)
qrows = [
    ('Critério ELS-D', '', 'CF: σ ≤ 0', 'CQP: σ ≤ 0', 'não exigido', None),
    ('Critério ELS-F', '', 'CR: σ ≤ 1,2·fctk,inf', 'CF: σ ≤ 1,2·fctk,inf', 'não exigido', None),
    ('Critério ELS-W', '', 'não exigido', 'não exigido', 'CF: wk ≤ 0,2 mm', None),
    ('Cordoalhas por cabo (2 cabos)', '', 5, 4, 2, 'in'),
    ('Cordoalhas totais Ø12,7', '', '=2*C8', '=2*D8', '=2*E8', '0'),
    ('Ap', 'cm²', f'=C9*{Dd("Ap1")}', f'=D9*{Dd("Ap1")}', f'=E9*{Dd("Ap1")}', '0.00'),
    ('Pj = 0,74·fptk·Ap', 'kN', f'=0.74*{Dd("fptk")}*C10/10', f'=0.74*{Dd("fptk")}*D10/10', f'=0.74*{Dd("fptk")}*E10/10', '0'),
    ('P∞ mín (após todas as perdas)', 'kN', 939, 776, 411, 'v'),
    ('P∞ máx', 'kN', 1050, 853, 442, 'v'),
    ('Carga equilibrada 8·f·P∞,mín/L²', 'kN/m', f'=8*({Dd("e_span")}+{Dd("e_sup")}/2)*C12/{Dd("L")}^2', f'=8*({Dd("e_span")}+{Dd("e_sup")}/2)*D12/{Dd("L")}^2', f'=8*({Dd("e_span")}+{Dd("e_sup")}/2)*E12/{Dd("L")}^2', '0.0'),
    ('Fração de g+ψ2q equilibrada', '%', f'=C14/{Dd("pbal")}*100', f'=D14/{Dd("pbal")}*100', f'=E14/{Dd("pbal")}*100', '0'),
    ('Margem ELS-D do critério do nível', 'MPa', 0.45, 1.04, 'n/a', 'v'),
    ('Margem ELS-F do critério do nível', 'MPa', 0.60, 3.14, 'n/a', 'v'),
    ('Atende ao nível 3?', '', 'sim', 'não (ELS-F rara: -1,49 MPa)', 'não', None),
    ('Atende ao nível 2?', '', 'sim', 'sim', 'não (ELS-D CQP)', None),
    ('Tração máx. na comb. frequente', 'MPa', -0.45, -0.70, 3.94, 'v'),
    ('Armadura passiva do ato (σs = 250 MPa)', 'cm²', 4.9, 3.2, 0.0, 'v'),
    ('Armadura passiva tracionada no apoio B (ELU)', 'cm²', '— (A\' = 4,9)', '2 Ø16 = 4,0', '4 Ø20 = 12,6', None),
    ('Armadura passiva tracionada no vão (ELU)', 'cm²', '—', 'não necessária', '2 Ø20 = 6,3', None),
    ('MRd/Msd no apoio B', '', 1.26, 1.14, 1.01, 'v'),
    ('MRd/Msd no vão', '', 1.60, 1.31, 1.09, 'v'),
    ('wk na comb. frequente (apoio B)', 'mm', 'não fissura', 'não fissura', 0.03, None),
    ('Flecha final (qp, com fluência; + para baixo)', 'mm', -2.8, -0.2, 5.4, 'v'),
    ('Massa de aço ativo (cabo de 20,09 m, 0,793 kg/m por cordoalha)', 'kg', '=C9*20.09*0.793', '=D9*20.09*0.793', '=E9*20.09*0.793', '0'),
    ('P∞ em x = 4 m', 'kN', 942, 779, 416, 'v0'),
    ('P∞ no apoio central B', 'kN', 976, 806, 431, 'v0'),
    ('Diâmetro externo da bainha (nível 2 e 1: adotado)', 'mm', 61, 55, 45, 'v0'),
    ('Alongamento por extremidade (depende só do atrito)', 'mm', 67.4, 67.4, 67.4, 'v1'),
    ('Desenho DXF 1:50 (A2)', '', 'Viga_Continua_CP_Forma_e_Cabos.dxf', 'Viga_Continua_CP_Nivel2_Forma_e_Cabos.dxf', 'Viga_Continua_CP_Nivel1_Forma_e_Cabos.dxf', None),
    ('Tração máx. na comb. quase permanente', 'MPa', -0.45, -1.04, 3.17, 'v'),
]
for k_, (lab, un, a_, b_, c3, fmt) in enumerate(qrows):
    rrw = 5 + k_
    wsQ.cell(rrw, 1, lab).font = fb; wsQ.cell(rrw, 2, un).font = fb
    for j, v in enumerate((a_, b_, c3), start=3):
        cell = wsQ.cell(rrw, j, v); cell.font = fb; cell.border = box
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        if fmt == 'in': cell.fill = fill_in
        elif fmt not in (None,): cell.fill = fill_res
        if fmt in ('0', '0.00', '0.0'): cell.number_format = fmt
wsQ.column_dimensions['A'].width = 56; wsQ.column_dimensions['B'].width = 8
for cc in 'CDE': wsQ.column_dimensions[cc].width = 26
wsQ.row_dimensions[4].height = 28

OUT = '/tmp/claude-0/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/scratchpad/Planilha_Apoio_Viga_Continua_CP.xlsx'
if __name__ == '__main__':
    wb.save(OUT)
    print('saved', OUT, 'Dados addresses sample', A['xa'], A['sfa'])
