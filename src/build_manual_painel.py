# -*- coding: utf-8 -*-
"""Gera o manual do painel (viga protendida de 1 vão e de 2 vãos): Word (.docx) e PDF (Chromium).
Uso, a partir da raiz do repositório:  py src/shots_painel.py  (telas)  e depois  py src/build_manual_painel.py"""
import os, re, struct, sys, html
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content_painel import blocos, VERSAO
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, 'src', 'manual_painel', 'img')
LOGO = os.path.join(ROOT, 'app', 'assets', 'logo-unb.png')
OUT = os.path.join(ROOT, 'entregas')
NOME = 'Manual_Painel_Viga_Protendida_1_e_2_vaos'
NAVY, BLUE, TINT = '164194', '0F6AB4', 'E3EAF6'
CALL = {'info': ('E3EAF6', '164194'), 'warn': ('FFF3D6', 'C98500'), 'ok': ('E3F4E3', '1E9B5B')}
W = 25.7                       # largura útil (cm), A4 paisagem com margens de 2 cm
DATA = '20/10/2026'

def dim(nome):
    with open(os.path.join(IMG, nome + '.png'), 'rb') as f: d = f.read(24)
    w, h = struct.unpack('>II', d[16:24]); return w / 1.5, h / 1.5     # px CSS (as telas foram feitas a 1,5x)

def fit(nome, maxw, maxh):
    w, h = dim(nome); k = min(maxw / w, maxh / h); return w * k, h * k  # cm, mantendo a proporção

# ====================================================================== Word
d = Document()
s = d.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE; s.page_width = Cm(29.7); s.page_height = Cm(21)
s.left_margin = s.right_margin = Cm(2); s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.5)
st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.12
for n, sz, col in (('Heading 1', 20, NAVY), ('Heading 2', 13.5, BLUE)):
    h = d.styles[n]; h.font.name = 'Arial'; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor.from_string(col)
    h.element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(6); h.paragraph_format.keep_with_next = True

def shade(c, color):
    p = c._tc.get_or_add_tcPr(); e = OxmlElement('w:shd'); e.set(qn('w:val'), 'clear'); e.set(qn('w:fill'), color); p.append(e)
def borders(c, color='C9D3E0', left=None, none=False):
    p = c._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for k in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + k)
        if none: e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '24' if (k == 'left' and left) else '4'); e.set(qn('w:color'), left if (k == 'left' and left) else color)
        b.append(e)
    p.append(b)
def runs(p, text, size=None, color=None, italic=False):
    for i, part in enumerate(re.split(r'(\*\*.+?\*\*|`.+?`)', text)):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'): r = p.add_run(part[1:-1]); r.font.name = 'Consolas'
        else: r = p.add_run(part)
        r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
def para(text='', size=None, color=None, bold=False, italic=False, align=None, after=None, container=None):
    p = (container or d).add_paragraph()
    if text: runs(p, text, size, color, italic);
    if bold:
        for r in p.runs: r.bold = True
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def spacer(pt=4):
    sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = Pt(pt); sp.add_run().font.size = Pt(2)
def fld(run, code):
    for t, x in (('begin', None), (None, code), ('end', None)):
        if t: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t)
        else: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = x
        run._r.append(e)
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run('Manual do painel · viga protendida de 1 vão e de 2 vãos · ' + VERSAO + ' · pág. '); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
r2 = fp.add_run(); r2.font.size = Pt(8.5); fld(r2, 'PAGE')

def add_pic(cell_or_doc, nome, wcm, hcm):
    p = cell_or_doc.add_paragraph() if hasattr(cell_or_doc, 'add_paragraph') else d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(IMG, nome + '.png'), width=Cm(wcm)); return p

def cell_clear(c):
    for p in list(c.paragraphs)[1:]: p._element.getparent().remove(p._element)
    c.paragraphs[0].text = ''
    return c.paragraphs[0]

def b_img(items, maxh):
    n = len(items); slot = (W - 0.4) / n
    t = d.add_table(rows=1, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (nome, leg) in enumerate(items):
        c = t.cell(0, j); c.width = Cm(W / n); borders(c, none=True); cell_clear(c)
        wcm, hcm = fit(nome, slot - 0.3, maxh); add_pic(c, nome, wcm, hcm)
        pl = c.add_paragraph(); pl.alignment = WD_ALIGN_PARAGRAPH.CENTER; pl.paragraph_format.space_after = Pt(0); runs(pl, leg, 8.5, '555555', True)
    spacer()
def b_side(nome, leg, maxh, textos):
    wcm, hcm = fit(nome, 12.5, maxh)
    t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.cell(0, 0), t.cell(0, 1); a.width = Cm(wcm + 0.6); b.width = Cm(W - wcm - 0.6); borders(a, none=True); borders(b, none=True)
    cell_clear(a); add_pic(a, nome, wcm, hcm)
    pl = a.add_paragraph(); pl.alignment = WD_ALIGN_PARAGRAPH.CENTER; pl.paragraph_format.space_after = Pt(0); runs(pl, leg, 8.5, '555555', True)
    p0 = cell_clear(b)
    for i, tx in enumerate(textos):
        p = p0 if i == 0 else b.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.4); p.paragraph_format.space_after = Pt(4)
        runs(p, '•  ' + tx)
    spacer()
def b_tab(rows, widths):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); borders(c); p = cell_clear(c); p.paragraph_format.space_after = Pt(1)
            runs(p, v, 9.5)
            if i == 0:
                for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
                shade(c, NAVY)
            elif i % 2 == 0: shade(c, 'F6F8FB')
    spacer()
def b_call(tipo, rot, texto):
    fill, bar = CALL[tipo]; t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); c.width = Cm(W); shade(c, fill); borders(c, 'FFFFFF', left=bar)
    p = cell_clear(c); p.paragraph_format.space_after = Pt(2); runs(p, '**' + rot + '** ' + texto); spacer()

def capa():
    para('', after=2)
    d.add_paragraph().add_run().add_picture(LOGO, width=Cm(5.2))
    para('UNIVERSIDADE DE BRASÍLIA · CONCRETO ARMADO E PROTENDIDO', size=10, color='8E2F6A', bold=True, after=4)
    para('Manual do painel interativo', size=30, color=NAVY, bold=True, after=2)
    para('Viga protendida de 1 vão e de 2 vãos', size=17, color=BLUE, after=10)
    para('Como usar o painel passo a passo, aba por aba, com as telas e os valores de conferência dos dois casos. Complemento da planilha de apoio, do memorial e do roteiro da disciplina.', size=11, after=14)
    para('Produzido pelo Professor DSc. Li Chong Lee (Todos os direitos reservados)', size=10, after=2)
    para('Painel ' + VERSAO + ' · ' + DATA, size=10, color='555555', after=2)
    para('https://lichonglee.github.io/viga-continua-protendida/app/viga_continua_protendida.html', size=9.5, color=BLUE, after=14)
    d.add_heading('Conteúdo', level=2)

# ====================================================================== HTML/PDF
H = []
def hesc(t):
    t = html.escape(t); t = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', t); return re.sub(r'`(.+?)`', r'<code>\1</code>', t)
def h_img(items, maxh):
    n = len(items); slot = (W - 0.4) / n; out = '<div class="row">'
    for nome, leg in items:
        wcm, hcm = fit(nome, slot - 0.3, maxh)
        out += '<figure style="width:%.2fcm"><img src="img/%s.png" style="width:%.2fcm"><figcaption>%s</figcaption></figure>' % (slot, nome, wcm, hesc(leg))
    H.append(out + '</div>')
def h_side(nome, leg, maxh, textos):
    wcm, hcm = fit(nome, 12.5, maxh)
    H.append('<div class="side"><figure style="width:%.2fcm"><img src="img/%s.png" style="width:%.2fcm"><figcaption>%s</figcaption></figure><ul>%s</ul></div>' % (wcm + .4, nome, wcm, hesc(leg), ''.join('<li>%s</li>' % hesc(t) for t in textos)))
def h_tab(rows, widths):
    o = '<table>'
    for i, row in enumerate(rows):
        o += '<tr>' + ''.join(('<th style="width:%.1fcm">%s</th>' if i == 0 else '<td>%s</td>') % ((widths[j], hesc(v)) if i == 0 else (hesc(v),)) for j, v in enumerate(row)) + '</tr>'
    H.append(o + '</table>')

# ====================================================================== monta
B = blocos(); caps = [b[1] for b in B if b[0] == 'h1']
capa()
for c in caps: para('•  ' + c, size=10.5, after=1)
H.append('<section class="capa"><img src="../../app/assets/logo-unb.png" class="logo"><p class="kick">UNIVERSIDADE DE BRASÍLIA · CONCRETO ARMADO E PROTENDIDO</p><h1 class="tit">Manual do painel interativo</h1><p class="sub">Viga protendida de 1 vão e de 2 vãos</p>'
         '<p>Como usar o painel passo a passo, aba por aba, com as telas e os valores de conferência dos dois casos. Complemento da planilha de apoio, do memorial e do roteiro da disciplina.</p>'
         '<p>Produzido pelo Professor DSc. Li Chong Lee (Todos os direitos reservados)</p><p class="mut">Painel %s · %s</p><p class="mut">https://lichonglee.github.io/viga-continua-protendida/app/viga_continua_protendida.html</p><h3>Conteúdo</h3><ul>%s</ul></section>' % (VERSAO, DATA, ''.join('<li>%s</li>' % hesc(c) for c in caps)))
numc = 0
for b in B:
    k = b[0]
    if k == 'h1':
        h = d.add_heading(b[1], level=1); h.paragraph_format.page_break_before = True; H.append('<h1>%s</h1>' % hesc(b[1]))
    elif k == 'h2': d.add_heading(b[1], level=2); H.append('<h2>%s</h2>' % hesc(b[1]))
    elif k == 'p': para(b[1]); H.append('<p>%s</p>' % hesc(b[1]))
    elif k == 'bul':
        for t in b[1]:
            p = d.add_paragraph(style='List Bullet'); runs(p, t); p.paragraph_format.space_after = Pt(2)
        H.append('<ul>%s</ul>' % ''.join('<li>%s</li>' % hesc(t) for t in b[1]))
    elif k == 'num':
        for i, t in enumerate(b[1], 1):
            p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.6); p.paragraph_format.space_after = Pt(2); runs(p, '%d.  %s' % (i, t))
        H.append('<ol>%s</ol>' % ''.join('<li>%s</li>' % hesc(t) for t in b[1]))
    elif k == 'img': b_img(b[1], b[2]); h_img(b[1], b[2])
    elif k == 'side': b_side(b[1], b[2], b[3], b[4]); h_side(b[1], b[2], b[3], b[4])
    elif k == 'tab': b_tab(b[1], b[2]); h_tab(b[1], b[2])
    elif k == 'call': b_call(b[1], b[2], b[3]); H.append('<div class="call %s"><b>%s</b> %s</div>' % (b[1], hesc(b[2]), hesc(b[3])))
os.makedirs(OUT, exist_ok=True)
docx_path = os.path.join(OUT, NOME + '.docx'); d.save(docx_path); print('docx', docx_path)

CSS = """@page{size:A4 landscape;margin:16mm 20mm 16mm 20mm}
body{font-family:Arial,Helvetica,sans-serif;font-size:10.5pt;line-height:1.42;color:#0B0B0B}
h1{font-size:20pt;color:#164194;border-bottom:3px solid #0095DB;padding-bottom:4px;margin:0 0 10px;page-break-before:always}
h2{font-size:13.5pt;color:#0F6AB4;margin:16px 0 6px;page-break-after:avoid}h3{font-size:12pt;color:#164194}
p{margin:0 0 7px}code{font-family:Consolas,monospace;font-size:9.5pt;background:#F1F3F7;padding:0 3px}
ul,ol{margin:2px 0 8px 22px;padding:0}li{margin:0 0 3px}
.capa{page-break-after:always}.capa .logo{width:5.2cm}.capa .kick{color:#8E2F6A;font-weight:700;font-size:10pt;margin-top:10px}
.capa .tit{font-size:30pt;margin:0;border:0;page-break-before:avoid}.capa .sub{font-size:17pt;color:#0F6AB4;margin-bottom:12px}.mut{color:#555}
.row{display:flex;gap:.4cm;justify-content:center;align-items:flex-start;margin:6px 0 10px;page-break-inside:avoid}
figure{margin:0;text-align:center}figure img{max-width:100%}figcaption{font-size:8.5pt;color:#555;font-style:italic;margin-top:2px}
.side{display:flex;gap:.6cm;align-items:flex-start;margin:6px 0 10px;page-break-inside:avoid}.side ul{flex:1}
table{border-collapse:collapse;width:100%;margin:6px 0 10px;font-size:9.5pt;page-break-inside:auto}th{background:#164194;color:#fff;text-align:left}
th,td{border:1px solid #C9D3E0;padding:3px 7px;vertical-align:top}tr:nth-child(even) td{background:#F6F8FB}tr{page-break-inside:avoid}
.call{border-left:5px solid;padding:6px 10px;margin:6px 0 10px;page-break-inside:avoid}.call.info{background:#E3EAF6;border-color:#164194}.call.warn{background:#FFF3D6;border-color:#C98500}.call.ok{background:#E3F4E3;border-color:#1E9B5B}"""
html_path = os.path.join(ROOT, 'src', 'manual_painel', 'manual.html')
open(html_path, 'w', encoding='utf-8').write('<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Manual do painel</title><style>%s</style></head><body>%s</body></html>' % (CSS, ''.join(H)))
from playwright.sync_api import sync_playwright
CHROME = os.path.expanduser(r'~\AppData\Local\ms-playwright\chromium-1248\chrome-win64\chrome.exe')
with sync_playwright() as p:
    br = p.chromium.launch(executable_path=CHROME); pg = br.new_page(); pg.goto('file:///' + html_path.replace('\\', '/')); pg.wait_for_timeout(800)
    pdf_path = os.path.join(OUT, NOME + '.pdf')
    pg.pdf(path=pdf_path, format='A4', landscape=True, print_background=True, display_header_footer=True, header_template='<span></span>',
           footer_template='<div style="font-size:8px;color:#666;width:100%;text-align:right;padding-right:20mm">Manual do painel · ' + VERSAO + ' · pág. <span class="pageNumber"></span> de <span class="totalPages"></span></div>',
           margin={'top': '16mm', 'bottom': '16mm', 'left': '20mm', 'right': '20mm'})
    br.close(); print('pdf', pdf_path)
