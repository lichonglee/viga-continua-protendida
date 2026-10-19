import sys,os
sys.path.insert(0,'.')
from content import ABAS
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
NAVY='154177';BLUE='0F6AB4';TINT='E8EEF5';WARN='FFF4D6';OK='E3F1E3'
d=Document()
s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width=Cm(29.7); s.page_height=Cm(21)
s.left_margin=s.right_margin=Cm(2); s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.6)
W=25.7
st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.12
for n,sz,col in (('Heading 1',20,NAVY),('Heading 2',14,BLUE)):
    h=d.styles[n]; h.font.name='Arial'; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=RGBColor.from_string(col)
    h.element.rPr.rFonts.set(qn('w:ascii'),'Arial'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    h.paragraph_format.space_before=Pt(10); h.paragraph_format.space_after=Pt(6); h.paragraph_format.keep_with_next=True
def shade(c,color):
    p=c._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:val'),'clear'); e.set(qn('w:fill'),color); p.append(e)
def borders(c,color='C9D3E0',left=None):
    p=c._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for k in ('top','left','bottom','right'):
        e=OxmlElement('w:'+k); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'18' if (k=='left' and left) else '4'); e.set(qn('w:color'),left if (k=='left' and left) else color); b.append(e)
    p.append(b)
def para(text='',bold=False,size=None,color=None,align=None,italic=False,after=None,style=None):
    p=d.add_paragraph(style=style); 
    if text:
        r=p.add_run(text); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor.from_string(color)
    if align: p.alignment=align
    if after is not None: p.paragraph_format.space_after=Pt(after)
    return p
def rich(p,label,text):
    r=p.add_run(label); r.bold=True; p.add_run(text)
def callout(label,text,fill,bar):
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); c.width=Cm(W); shade(c,fill); borders(c,'FFFFFF',left=bar)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); rich(p,label+' ',text)
    sp=d.add_paragraph(); sp.paragraph_format.space_after=Pt(0); sp.paragraph_format.line_spacing=Pt(4); sp.add_run().font.size=Pt(2)
def bullets(items):
    for it in items:
        p=d.add_paragraph(style='List Bullet'); p.add_run(it); p.paragraph_format.space_after=Pt(2)
def table(rows,widths,head=True,size=9.5):
    t=d.add_table(rows=len(rows),cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.width=Cm(widths[j]); borders(c)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
            r=p.add_run(v); r.font.size=Pt(size)
            if head and i==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); shade(c,NAVY)
            elif i%2==0: shade(c,'F6F8FB')
    sp=d.add_paragraph(); sp.paragraph_format.space_after=Pt(0); sp.paragraph_format.line_spacing=Pt(4); sp.add_run().font.size=Pt(2)
def fld(run,code):
    for t,x in (('begin',None),(None,code),('end',None)):
        if t: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
        else: e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=x
        run._r.append(e)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=fp.add_run('Manual da planilha de apoio · viga contínua protendida · pág. '); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x66,0x66,0x66)
r2=fp.add_run(); r2.font.size=Pt(8.5); fld(r2,'PAGE')

# capa
para('UNIVERSIDADE DE BRASÍLIA · CONCRETO ARMADO E PROTENDIDO',bold=True,size=10,color='8E2F6A',after=4)
p=para('Manual da planilha de apoio',bold=True,size=30,color=NAVY,after=2)
para('Viga contínua protendida de dois vãos · 14 abas, tela a tela',size=15,color=BLUE,after=8)
para('Complemento do "Roteiro para reproduzir as planilhas da viga contínua protendida". O roteiro diz como montar cada aba; este manual mostra como cada aba deve ficar na tela, o que olhar nela e que valores conferir antes de seguir para a próxima.',after=6)
para('Prof. Li Chong Lee Bacelar de Castro · Seção 30 × 90 cm · C30 · CAA II · CP190 RB Ø12,7 · 2 × 10 m',size=9.5,color='555555',after=10)
d.add_heading('Como usar este manual',level=2)
bullets(['Siga a ordem das abas, que é a do roteiro: cada uma consulta as anteriores, então um erro cedo contamina tudo o que vem depois.',
 'Em cada aba, compare sua tela com a figura, depois confira a caixa "Valores de conferência". Se um valor falhar, use a caixa "Erro comum" e volte ao passo indicado do roteiro.',
 'As telas mostram as primeiras linhas de cada aba (as tabelas por estação vão até a linha 48, x = 20 m) e, nas abas largas, as colunas em duas faixas.',
 'Os valores das telas são os da planilha de referência, com os dados do projeto descritos no roteiro.'])
d.add_heading('Convenções da planilha',level=2)
table([['Elemento','Significado'],
 ['Célula amarela','Entrada: pode ser alterada pelo usuário (dados, escolha de barras, número de cordoalhas).'],
 ['Célula verde ou azul-clara','Resultado: fórmula ou valor de conferência. Não digitar por cima.'],
 ['Cabeçalho azul-escuro','Nomes das colunas (linha 6) e unidades (linha 7) nas abas por estação.'],
 ['Estações','41 linhas, de x = 0 (linha 8) a x = 20 m (linha 48), a cada 0,5 m; B está na linha 28.'],
 ['Sinais','Tração positiva, compressão negativa; momento positivo = tração embaixo; e positivo abaixo do CG; σ1 = fibra inferior, σ2 = fibra superior.'],
 ['Unidades','kN, m e MPa; áreas de aço em cm²; aberturas de fissura e flechas em mm.']],[6,19.7])
hh=d.add_heading('Mapa das abas',level=2); hh.paragraph_format.page_break_before=True
rows=[['Nº','Aba','Roteiro','Para que serve','Depende de']]
fn={'Dados':'Entradas do projeto','Esforcos':'Momentos e cortantes externos','Tracado_Perdas':'Traçado do cabo e perdas','Hiperestatico':'Momento secundário da protensão','Tensoes':'Tensões no ato e em serviço','Fuso':'Fuso limite e linha de pressões','Cortante':'Estribos (Modelo I)','ELU':'Flexão na ruptura','Alongamento':'Alongamento dos cabos','Flecha_PC':'Flecha da viga protendida','CA':'Viga em concreto armado','CA_Flecha':'Flecha em concreto armado (Branson)','Niveis':'Níveis de protensão do projeto atual','Niveis_Comparativo':'Níveis 3, 2 e 1 lado a lado'}
dep={'Dados':'—','Esforcos':'Dados','Tracado_Perdas':'Dados','Hiperestatico':'Esforcos, Tracado_Perdas','Tensoes':'Esforcos, Tracado_Perdas, Hiperestatico','Fuso':'Tensoes, Hiperestatico','Cortante':'Esforcos, Tracado_Perdas, Hiperestatico','ELU':'Esforcos, Tracado_Perdas, Hiperestatico','Alongamento':'Tracado_Perdas','Flecha_PC':'Esforcos, Hiperestatico','CA':'Esforcos, Dados','CA_Flecha':'Esforcos, CA','Niveis':'Tensoes, Esforcos','Niveis_Comparativo':'Niveis (rodadas)'}
for i,a in enumerate(ABAS,1): rows.append([str(i),a[0],a[1],fn[a[0]],dep[a[0]]])
table(rows,[1.2,4.5,2.6,8.4,9])
for i,(n,passo,shots,obj,obs,conf,erro) in enumerate(ABAS,1):
    hh=d.add_heading(f'Aba {i}: {n}',level=1); hh.paragraph_format.page_break_before=True
    para(passo+' do roteiro',italic=True,size=9.5,color='666666',after=4)
    para(obj)
    if n=='Dados':
        t=d.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for j,(f,leg) in enumerate(shots):
            w,h=Image.open(f).size; c=t.cell(0,j); c.width=Cm(12.8)
            c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].add_run().add_picture(f,width=Cm(min(12.4,9.0*w/h)))
            q=t.cell(1,j).paragraphs[0]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run('Figura: '+leg); r.italic=True; r.font.size=Pt(8.5)
        shots=[]
    for f,leg in shots:
        w,h=Image.open(f).size; ww=min(W,w/ (h/ 9.5)) if False else None
        wd=W if w/h>1.8 else min(W,13.5*w/h*1.0) if h>w*0.6 else W
        maxh=8.4 if len(shots)==1 else 8.8
        wd=min(W,maxh*w/h)
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(f,width=Cm(wd))
        para('Figura: '+leg,italic=True,size=9,color='555555',align=WD_ALIGN_PARAGRAPH.CENTER)
    d.add_heading('O que observar',level=2); bullets(obs)
    callout('Valores de conferência:',conf,OK,'3A8F3A')
    callout('Erro comum:',erro,WARN,'F39325')
hh=d.add_heading('Conferência final em uma tela',level=1); hh.paragraph_format.page_break_before=True
table([['Aba','Grandeza','Valor de referência'],
['Dados','Ap; p_bal; e_max; Pj','9,87 cm²; 27,5 kN/m; 37,45 cm; 1388 kN'],
['Esforcos','Mg e Mq,mín em B','−250,0 e −312,5 kN·m'],
['Tracado_Perdas','xa; Σα; comprimento do cabo','9,23 m; 1,013 rad; 20,09 m'],
['Tracado_Perdas','P∞ em x = 4 m e em B','941,7 kN e 975,7 kN'],
['Hiperestatico','R_B e M_hip,B em t∞','−13,78 kN e +68,9 kN·m'],
['Tensoes','σ1 máx no ato; As,ato','3,90 MPa (ARMAR em B); 4,9 cm²'],
['Tensoes','CFQ máx; CR máx','−0,45 MPa; 1,83 MPa (limite 2,43)'],
['Cortante','Vsd máx; VRd2','388 kN; 986 kN'],
['ELU','MRd/Msd em B e no vão','1,26 e 1,60'],
['Alongamento','por extremidade; total; líquido','67,4 mm; 134,8 mm; 61,4 mm por extremidade'],
['Flecha_PC','δ_qp,∞ mínimo','−2,8 mm (contraflecha)'],
['CA','As,calc em B e no vão; wk','25,2 e 16,1 cm² (adotados 8 Ø20 e 6 Ø20); 0,10 mm'],
['CA_Flecha','δ total','+7,4 mm'],
['Niveis','com C57 = 5, 4 e 2','5: níveis 3 e 2; 4: nível 2; 2: só o nível 1']],[4.5,8,13.2])
d.add_heading('Dicas de navegação e de trabalho',level=2)
bullets(['Ctrl + Page Down e Ctrl + Page Up passam de uma aba para a seguinte e a anterior.',
 'Congele os painéis na linha 8 das abas por estação (Exibir > Congelar painéis) para manter os cabeçalhos visíveis.',
 'Ctrl + ` (acento grave) alterna entre ver valores e ver fórmulas, útil para comparar com o roteiro.',
 'Se aparecer #NOME?, o idioma do Excel não reconhece o nome da função: troque SE/E/MÁXIMO/RAIZ por IF/AND/MAX/SQRT ou o inverso.',
 'Fórmulas que citam a aba Dados devem usar endereço absoluto (por exemplo Dados!$C$5), senão a cópia para baixo desloca a referência.',
 'Ao terminar, proteja as células de fórmula e deixe editáveis só as amarelas.'])
d.save('Manual_Planilha_Apoio_Viga_Continua_CP.docx')
