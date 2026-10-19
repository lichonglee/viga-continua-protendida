import json, re, sys, os, hashlib, subprocess, shutil
import xml.etree.ElementTree as ET
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE=os.path.dirname(os.path.abspath(__file__))
TR='/root/.claude/projects/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/tool-results/'
NAVY='154177'; BLUE='0F6AB4'; ORANGE='F39325'; TINT='E8EEF5'; WARN='FFF4D6'; GREY='F2F2F2'
rgb=lambda h: RGBColor.from_string(h)
os.makedirs(HERE+'/fx',exist_ok=True)

# ---------------------------------------------------------------- latex -> png
NONASCII=re.compile(r'[A-Za-z]*[^\x00-\x7f][A-Za-zÀ-ſ]*')
def prep(tex):
    tex=tex.replace('‰',r'\text{\textperthousand}')
    parts=re.split(r'(\\text\{[^{}]*\})',tex)
    out=[]
    for p in parts:
        if p.startswith('\\text{'): out.append(p)
        else: out.append(NONASCII.sub(lambda m:'\\text{'+m.group(0)+'}',p))
    return ''.join(out)
def latex_png(tex):
    h=hashlib.md5(tex.encode()).hexdigest()[:12]; png=f'{HERE}/fx/{h}.png'
    if os.path.exists(png): return png
    src=r'''\documentclass[11pt,preview,border=4pt]{standalone}
\usepackage[utf8]{inputenc}\usepackage{amsmath,amssymb,textcomp}
\begin{document}
\begin{math}\displaystyle %s\end{math}
\end{document}'''%prep(tex)
    d=f'{HERE}/fx/{h}'; os.makedirs(d,exist_ok=True)
    open(d+'/f.tex','w').write(src)
    r=subprocess.run(['pdflatex','-interaction=nonstopmode','-halt-on-error','f.tex'],cwd=d,capture_output=True,text=True)
    if not os.path.exists(d+'/f.pdf'):
        print('LATEX FAIL',tex[:80],r.stdout[-400:]); return None
    subprocess.run(['pdftoppm','-r','300','-png','-singlefile','f.pdf','f'],cwd=d)
    shutil.move(d+'/f.png',png); return png

# ---------------------------------------------------------------- xml -> blocks
def runs_of(el, style=()):
    out=[]
    if el.text: out.append((el.text,style))
    for ch in el:
        tag=ch.tag
        if tag in ('bold','italic','code','link'):
            out+=runs_of(ch, style+((tag,),) if False else style+(tag,))
        elif tag=='text':
            out+=runs_of(ch,style)
        if ch.tail: out.append((ch.tail,style))
    return out
def para_runs(p):
    out=[]
    for ch in p:
        if ch.tag=='text': out+=runs_of(ch)
    return out
def parse(path):
    j=json.load(open(path)); x=j['data']['xml']
    root=ET.fromstring(x); blocks=[]
    def cell_text(c):
        return [para_runs(p) for p in c if p.tag=='paragraph']
    for el in root:
        t=el.tag
        if t=='paragraph':
            rs=para_runs(el)
            if el.get('heading'): blocks.append(('h',int(el.get('heading')),rs))
            elif any(c.tag in('date','mention') for c in el): blocks.append(('byline',))
            else: blocks.append(('p',rs))
        elif t=='list':
            items=[]
            for li in el:
                ps=[para_runs(p) for p in li if p.tag=='paragraph']
                items.append(ps[0] if ps else [])
            blocks.append(('list',el.get('kind'),int(el.get('start','1')),items))
        elif t=='codeBlock':
            blocks.append(('code',el.get('language'),''.join(el.itertext())))
        elif t=='table':
            rows=[]
            for r in el:
                rows.append([ (c.get('header')=='true', cell_text(c)) for c in r])
            blocks.append(('table',rows))
        elif t=='embed': blocks.append(('embed',el.get('caption')))
    return blocks

# ---------------------------------------------------------------- docx helpers
def shade(cell_or_p,color,is_p=False):
    el=cell_or_p._p.get_or_add_pPr() if is_p else cell_or_p._tc.get_or_add_tcPr()
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),color); el.append(sh)
def borders(cell,**kw):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        e=OxmlElement('w:'+side)
        if side in kw: v=kw[side]; e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(v[0])); e.set(qn('w:color'),v[1])
        else: e.set(qn('w:val'),'nil')
        b.append(e)
    tcPr.append(b)
def pborder(p,side,sz,color,space=1):
    pPr=p._p.get_or_add_pPr(); b=pPr.find(qn('w:pBdr'))
    if b is None: b=OxmlElement('w:pBdr'); pPr.append(b)
    e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz)); e.set(qn('w:space'),str(space)); e.set(qn('w:color'),color); b.append(e)
def field(run,code):
    for typ,txt in (('begin',None),(None,code),('end',None)):
        if typ:
            f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),typ); run._r.append(f)
        else:
            it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=txt; run._r.append(it)
def add_runs(p,rs,size=None,color=None,bold_all=False):
    for t,st in rs:
        r=p.add_run(t)
        if 'bold' in st or bold_all: r.bold=True
        if 'italic' in st: r.italic=True
        if 'code' in st:
            r.font.name='Consolas'; r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Consolas'); r.font.size=Pt((size or 10.5)-1); r.font.color.rgb=rgb('922F5F')
        else:
            if size: r.font.size=Pt(size)
            if color: r.font.color.rgb=rgb(color)
def plain(rs): return ''.join(t for t,_ in rs)

CALL=('Confira','Atenção','Verificação obrigatória','Leitura:','Pendências','Dica','Como trabalhar','Tolerância')
def build(blocks,out,title,subtitle,fig_map):
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    pf=st.paragraph_format; pf.space_after=Pt(5); pf.line_spacing=1.1
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'pt-BR'); st.element.rPr.append(lang)
    for n,sz,col in (('Heading 1',15,NAVY),('Heading 2',12.5,BLUE),('Heading 3',11,NAVY)):
        s=doc.styles[n]; s.font.name='Calibri'; s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=rgb(col)
        s.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri'); s.element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s.element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
        s.paragraph_format.space_before=Pt(16 if n=='Heading 1' else 10); s.paragraph_format.space_after=Pt(6); s.paragraph_format.keep_with_next=True
    # ---- capa
    sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(0); sec.top_margin=Cm(0); sec.bottom_margin=Cm(2.4)
    sec.header_distance=Cm(0); sec.footer_distance=Cm(0)
    sec.footer.is_linked_to_previous=False
    fq=sec.footer.paragraphs[0]; fq.paragraph_format.space_after=Pt(0); fq.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fq.add_run().add_picture(HERE+'/band.png',width=Cm(21))
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(HERE+'/banner_s.png',width=Cm(21))
    def cp(text,size,bold=False,color=NAVY,after=6,italic=False,before=0):
        q=doc.add_paragraph(); q.paragraph_format.left_indent=Cm(2.2); q.paragraph_format.right_indent=Cm(2.2)
        q.paragraph_format.space_after=Pt(after); q.paragraph_format.space_before=Pt(before)
        r=q.add_run(text); r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=rgb(color); return q
    cp('ROTEIRO DE ATIVIDADE',11,True,ORANGE,after=4,before=40)
    cp(title,26,True,NAVY,after=10)
    q=cp('',1,after=14); pborder(q,'bottom',18,ORANGE,1)
    cp(subtitle,13,False,'404040',after=24)
    cp('Prof. Li Chong Lee Bacelar de Castro',12,True,NAVY,after=2,before=30)
    cp('Concreto Protendido · Universidade de Brasília',11,False,'404040',after=2)
    cp('Brasília, outubro de 2026',11,False,'404040',after=2)
    # ---- corpo
    sec2=doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.left_margin=sec2.right_margin=Cm(2.2); sec2.top_margin=Cm(2.4); sec2.bottom_margin=Cm(2.2)
    sec2.header_distance=Cm(1.0); sec2.footer_distance=Cm(0.8)
    sec2.header.is_linked_to_previous=False; sec2.footer.is_linked_to_previous=False
    hp=sec2.header.paragraphs[0]; hp.text=''
    r=hp.add_run('UNIVERSIDADE DE BRASÍLIA'); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=rgb(NAVY)
    r=hp.add_run('  ·  Concreto Armado e Protendido'); r.font.size=Pt(8.5); r.font.color.rgb=rgb('595959')
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.2),WD_TAB_ALIGNMENT.RIGHT)
    r=hp.add_run('\t\tRoteiro de atividade'); r.font.size=Pt(8.5); r.font.color.rgb=rgb('595959')
    pborder(hp,'bottom',12,ORANGE,4)
    fp=sec2.footer.paragraphs[0]; fp.text=''
    shade(fp,NAVY,True); fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.2),WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.left_indent=Cm(0.2); fp.paragraph_format.right_indent=Cm(0.2)
    r=fp.add_run('Prof. Li Chong Lee Bacelar de Castro'); r.font.size=Pt(8.5); r.font.color.rgb=rgb('FFFFFF')
    r=fp.add_run('\t\tPágina '); r.font.size=Pt(8.5); r.font.color.rgb=rgb('FFFFFF')
    r=fp.add_run(); r.font.size=Pt(8.5); r.font.color.rgb=rgb('FFFFFF'); field(r,'PAGE')
    r=fp.add_run(' de '); r.font.size=Pt(8.5); r.font.color.rgb=rgb('FFFFFF')
    r=fp.add_run(); r.font.size=Pt(8.5); r.font.color.rgb=rgb('FFFFFF'); field(r,'NUMPAGES')

    def callout(rs,warn=False):
        tb=doc.add_table(rows=1,cols=1); tb.alignment=WD_TABLE_ALIGNMENT.CENTER
        c=tb.rows[0].cells[0]; shade(c,WARN if warn else TINT); borders(c,left=(36,ORANGE if warn else BLUE))
        q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(2); q.paragraph_format.space_before=Pt(2)
        q.paragraph_format.left_indent=Cm(0.15)
        add_runs(q,rs,size=10)
        sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(2); sp.paragraph_format.line_spacing=0.6
    first_h1=True; figno=0
    for b in blocks:
        k=b[0]
        if k=='byline': continue
        if k=='h':
            lvl,rs=b[1],b[2]
            if lvl==1: continue          # título vai para a capa
            doc.add_heading(plain(rs),level=lvl-1 if lvl>1 else 1); continue
        if k=='p':
            rs=b[1]; txt=plain(rs).strip()
            if not txt: continue
            if txt.startswith(CALL):
                callout(rs,warn=txt.startswith(('Atenção','Pendências','Verificação')))
            else:
                p=doc.add_paragraph(); add_runs(p,rs)
            continue
        if k=='list':
            kind,start,items=b[1],b[2],b[3]
            for i,rs in enumerate(items):
                p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(1.0); p.paragraph_format.first_line_indent=Cm(-0.55)
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
                mark={'ordered':f'{start+i}.','check':'☐'}.get(kind,'•')
                r=p.add_run(mark+'\t'); r.bold=True; r.font.color.rgb=rgb(BLUE if kind!='ordered' else NAVY)
                add_runs(p,rs)
            continue
        if k=='code':
            lang,src=b[1],b[2]
            if lang=='latex':
                png=latex_png(src.strip())
                if png:
                    w=Image.open(png).size[0]/300*2.54*1.0
                    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(6)
                    p.paragraph_format.keep_together=True
                    p.add_run().add_picture(png,width=Cm(min(w,16.6)))
                    continue
            tb=doc.add_table(rows=1,cols=1); c=tb.rows[0].cells[0]; shade(c,GREY); borders(c,left=(24,NAVY))
            c.paragraphs[0].text=''
            for i,l in enumerate(src.split('\n')):
                q=c.paragraphs[0] if i==0 else c.add_paragraph(); q.paragraph_format.space_after=Pt(0)
                r=q.add_run(l); r.font.name='Consolas'; r.font.size=Pt(9)
            doc.add_paragraph().paragraph_format.space_after=Pt(2); continue
        if k=='table':
            rows=b[1]; nc=max(len(r) for r in rows)
            t=doc.add_table(rows=0,cols=nc); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
            lens=[0]*nc
            for r in rows:
                for j,(h,ps) in enumerate(r):
                    lens[j]=max(lens[j],min(40,max([len(plain(x)) for x in ps] or [1])))
            tot=sum(max(l,6) for l in lens); W=16.6
            widths=[W*max(l,6)/tot for l in lens]
            fs=9.5 if nc<=4 else 8.5
            for ri,r in enumerate(rows):
                cells=t.add_row().cells
                hdr=all(h for h,_ in r)
                for j in range(nc):
                    c=cells[j]; c.width=Cm(widths[j])
                    ps=r[j][1] if j<len(r) else []
                    c.paragraphs[0].text=''
                    for pi,rs in enumerate(ps or [[]]):
                        q=c.paragraphs[0] if pi==0 else c.add_paragraph(); q.paragraph_format.space_after=Pt(1); q.paragraph_format.space_before=Pt(1); q.paragraph_format.line_spacing=1.0
                        add_runs(q,rs,size=fs,color='FFFFFF' if hdr else None,bold_all=hdr)
                    if hdr: shade(c,NAVY)
                    elif ri%2==0: shade(c,TINT)
                    borders(c,top=(4,'BFC9D6'),bottom=(4,'BFC9D6'),left=(4,'BFC9D6'),right=(4,'BFC9D6'))
                if hdr:
                    trPr=t.rows[-1]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); trPr.append(th)
                cs=OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(cs)
            doc.add_paragraph().paragraph_format.space_after=Pt(4); continue
        if k=='embed':
            cap=b[1]; png=fig_map.get(cap)
            if png:
                figno+=1
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(png,width=Cm(12.2))
                q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=q.add_run(f'Figura {figno} – '+cap); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=rgb('595959')
            continue
    doc.core_properties.title=title; doc.core_properties.author='Li Chong Lee Bacelar de Castro'
    doc.save(out)

if __name__=='__main__':
    Image.open('/tmp/claude-0/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/scratchpad/unbt/ppt/media/image1.png').convert('RGB').crop((0,860,1920,1080)).save(HERE+'/band.png')
    im=Image.open(HERE+'/banner.png'); im=im.resize((1800,int(1800*im.size[1]/im.size[0])),Image.LANCZOS); im.save(HERE+'/banner_s.png')
    OUT='/tmp/claude-0/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/scratchpad/out/'
    bpmn='/root/.claude/projects/-home-claude/714531a0-74fd-5a12-b94f-74c52a8d18b7/tool-results/mcp-Claude_Docs-blob-1791143071249-sq8p5y.png'
    sub='Viga contínua protendida de dois vãos de 10 m — seção 30 × 90 cm — C30 — CAA II — CP 190 RB Ø12,7 — pós-tração aderente'
    b2=parse(TR+'mcp-Claude_Docs-read-1791207777925.txt')
    build(b2,OUT+'Roteiro_Viga_Continua_Protendida_UnB.docx','Roteiro para reproduzir a viga contínua protendida de 2 vãos',sub,{'fluxograma BPMN do roteiro · 2 raias, 8 passos, 2 gateways':bpmn})
    b1=parse(TR+'mcp-Claude_Docs-read-1791207777892.txt')
    build(b1,OUT+'Roteiro_Planilhas_Viga_Continua_UnB.docx','Roteiro para reproduzir as planilhas da viga contínua protendida',sub,{})
    print(len(b1),len(b2))
